"""Build the two Caldova session decks (Microsoft-branded).

Deck 1 covers Demo 1 (local MCP + multi-agent engineering).
Deck 2 covers Demo 2 (Microsoft Foundry Hosted Agent production operation).

Both are aimed at developers and AI engineers, are technically detailed, and
include a dedicated live-demo placeholder slide. Run: `python build_deck.py`.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from deck_common import (
    AMBER, AMBER_SOFT, CYAN, DEEP, GREEN, GREEN_SOFT, INK, LINE, MSBLUE,
    MSBLUE_DK, MUTED, PURPLE, RED, RED_SOFT, SLIDE_H, SLIDE_W, WHITE,
    add_notes, bullet_row, closing_slide, code_line, demo_slide, heading,
    label, line, node, rect, text, title_slide,
)


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
DISCLAIMER = "Fictional pharmaceutical recall. Synthetic operational data. Batch B-2408-AX7."


def add_picture(slide, filename, x, y, w, h):
    return slide.shapes.add_picture(str(ASSETS / filename), Inches(x), Inches(y), Inches(w), Inches(h))


# ---------------------------------------------------------------------------
# Deck 1 — Demo 1: Local MCP + multi-agent engineering
# ---------------------------------------------------------------------------

def d1_title(prs):
    title_slide(
        prs,
        demo_tag="Demo 1 · Local engineering",
        title_value="Engineering the Agent",
        subtitle="MCP as a protocol contract, multi-agent least privilege, and deterministic authority — built and tested locally",
        chips=["MCP 2026-07-28", "mcp 2.1.0 SDK", "4 least-priv agents", "stdio + in-memory"],
        disclaimer=DISCLAIMER,
    )


def d1_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "The engineering shift", "A capable model is not a reliable system", 2)
    text(slide, "PROMPT LAYER", 0.62, 1.72, 3.0, 0.2, size=9, value_color=RED, bold=True)
    rect(slide, 0.62, 2.0, 3.7, 3.8, WHITE, line=LINE)
    text(slide, "The model can…", 0.9, 2.28, 3.0, 0.35, size=20, font="Segoe UI Semibold", bold=True)
    for index, item in enumerate(["Interpret the recall notice", "Summarize supplier evidence", "Recommend the next action"]):
        bullet_row(slide, item, 0.9, 2.95 + index * 0.62, 3.2, size=13, accent=RED)
    label(slide, "Behavior, not authority", 0.9, 5.15, 2.2, RED_SOFT, RED)

    text(slide, "SYSTEM LAYER", 4.78, 1.72, 3.0, 0.2, size=9, value_color=MSBLUE, bold=True)
    rect(slide, 4.78, 2.0, 7.9, 3.8, DEEP)
    text(slide, "The application must…", 5.1, 2.28, 5.0, 0.35, size=20, font="Segoe UI Semibold", value_color=WHITE, bold=True)
    controls = [
        ("IDENTITY", "Know which human or workload is acting"),
        ("POLICY", "Validate approval below the model prompt"),
        ("REPLAY", "Make every state mutation idempotent"),
        ("EVIDENCE", "Correlate model, workflow, tool and mutation"),
    ]
    for index, (name, detail) in enumerate(controls):
        x = 5.1 + (index % 2) * 3.75
        y = 3.0 + (index // 2) * 1.15
        text(slide, name, x, y, 1.6, 0.2, size=9, value_color=CYAN, bold=True)
        text(slide, detail, x, y + 0.26, 3.2, 0.6, size=12.5, value_color="E4ECF5")
    text(slide, "Prompts guide behavior. Deterministic controls grant authority.",
         4.78, 6.15, 7.9, 0.4, size=17, font="Segoe UI Semibold", value_color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, """
[Framing for engineers] Prompt instructions are useful behavior guidance, but they are not identity, authorization, validation, or replay protection. Everything in this deck pushes authority below the model. Ask the room: if the same tool call is retried after a timeout, what prevents a second mutation?
""")


def d1_scenario(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "The scenario as a test contract", "One batch. Four locations. One controlled decision.", 3)
    label(slide, "High-risk recall", 0.58, 1.7, 1.6, RED)
    text(slide, "Caldova Relief 20 mg tablets", 0.58, 2.1, 6.2, 0.48, size=27, font="Segoe UI Semibold", bold=True)
    text(slide, "Batch B-2408-AX7  ·  Northstar Therapeutics  ·  cold-chain excursion", 0.58, 2.64, 6.1, 0.3, size=13, value_color=MUTED)
    metrics = [("2,196", "affected units", RED), ("4", "network locations", MSBLUE), ("1", "named approval", AMBER)]
    for index, (value, caption, accent) in enumerate(metrics):
        x = 0.58 + index * 2.07
        rect(slide, x, 3.2, 1.78, 1.42, WHITE, line=LINE)
        rect(slide, x, 3.2, 1.78, 0.07, accent)
        text(slide, value, x + 0.16, 3.46, 1.5, 0.48, size=29, value_color=accent, font="Segoe UI Semibold", bold=True)
        text(slide, caption, x + 0.16, 4.0, 1.5, 0.22, size=10, value_color=MUTED)
    text(slide, "Fixed invariants make both demos repeatable and assertable in tests.",
         0.58, 4.85, 6.2, 0.5, size=12.5, value_color=INK)
    for index, loc in enumerate(["Bengaluru DC — 1,240", "Chennai DC — 760", "Bengaluru Store 031 — 112", "Mysuru Store 014 — 84"]):
        bullet_row(slide, loc, 0.58, 5.45 + index * 0.36, 6.2, size=11, accent=MSBLUE, font="Cascadia Code")

    rect(slide, 7.05, 1.7, 5.65, 4.9, DEEP)
    text(slide, "DEFINITION OF DONE", 7.38, 2.0, 3.0, 0.22, size=9, value_color=CYAN, bold=True)
    for index, item in enumerate([
        "Establish recall facts from typed tools",
        "Compute exposure from returned positions",
        "Require named human approval",
        "Quarantine every position exactly once",
        "Return sanitized, correlated evidence",
    ]):
        y = 2.5 + index * 0.72
        rect(slide, 7.38, y, 0.24, 0.24, MSBLUE)
        text(slide, str(index + 1), 7.38, y + 0.02, 0.24, 0.14, size=8, value_color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(slide, item, 7.82, y - 0.01, 4.6, 0.5, size=13.5, value_color=WHITE)
    add_notes(slide, "The five outcomes are the acceptance contract for BOTH demos. Every claim later maps back to one of these.")


def d1_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Local architecture", "Separate reasoning, protocol and authority", 4)
    node(slide, "Recall supervisor", "Named human decision", 0.55, 2.4, 1.8, 1.0, AMBER)
    node(slide, "Control Tower", "Operational UI + API", 2.75, 2.4, 1.8, 1.0, MSBLUE)
    node(slide, "Hosted workflow", "Responses protocol", 4.95, 2.4, 1.95, 1.0, MSBLUE)
    node(slide, "4 specialist agents", "Least-privilege reasoning", 7.3, 1.55, 2.1, 1.0, MSBLUE)
    node(slide, "Model (local run)", "Agent Framework client", 7.3, 3.26, 2.1, 1.0, MSBLUE)
    node(slide, "MCP v2 boundary", "Protocol 2026-07-28", 9.8, 2.4, 2.05, 1.0, GREEN)
    node(slide, "Policy + store", "Approval / idempotency / audit", 9.8, 4.25, 2.05, 1.0, RED)
    for x1, y1, x2, y2 in [
        (2.35, 2.9, 2.75, 2.9), (4.55, 2.9, 4.95, 2.9), (6.9, 2.9, 7.3, 2.05),
        (6.9, 2.9, 7.3, 3.76), (9.4, 2.05, 9.8, 2.9), (11.05, 3.4, 11.05, 4.25),
        (3.65, 3.4, 9.8, 4.75),
    ]:
        line(slide, x1, y1, x2, y2, MSBLUE if y2 < 4.0 else RED, 1.5, dash=y2 >= 4.0)
    text(slide, "read-only tools", 9.77, 1.92, 1.6, 0.2, size=9, value_color=GREEN, bold=True)
    text(slide, "mutate only after approval", 6.7, 5.16, 2.4, 0.2, size=9, value_color=RED, bold=True)
    rect(slide, 0.55, 6.17, 12.15, 0.54, GREEN_SOFT)
    text(slide, "The application owns authority; agents own bounded reasoning; MCP owns the tool contract.",
         0.78, 6.31, 11.68, 0.22, size=15, value_color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Call out the dashed red path: mutation never originates from an agent prompt. Demo 1 starts at the MCP boundary because that is where prototype ambiguity becomes a contract.")


def d1_mcp(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "MCP v2 engineering", "A protocol contract, not a bag of functions", 5)
    text(slide, "WHAT WE STANDARDISED", 0.58, 1.62, 5.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, item in enumerate([
        "MCPServer (SDK v2) — not the deprecated FastMCP v1 API",
        "Strict typed inputs → generated JSON Schema 2020-12",
        "Structured content plus backward-compatible text",
        "Per-request protocol metadata; no session handshake",
        "Built-in OpenTelemetry middleware; W3C trace context",
    ]):
        bullet_row(slide, item, 0.58, 2.0 + index * 0.5, 6.5, size=12.5, accent=MSBLUE)

    text(slide, "ISOLATION THAT MATTERS", 0.58, 4.6, 6.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    text(slide, "Foundry hosting needs mcp<2, but the demo needs mcp 2.1.0. The bridge runs the MCP v2 server in an isolated interpreter and calls it with the in-memory Client transport — no nested subprocess.",
         0.58, 4.92, 6.5, 1.1, size=12.5, value_color=INK)

    rect(slide, 7.35, 1.62, 5.35, 3.15, DEEP)
    text(slide, "VALIDATED CONTRACT", 7.65, 1.9, 3.0, 0.22, size=9, value_color=CYAN, bold=True)
    for index, (key, value) in enumerate([
        ("SDK", "mcp 2.1.0"), ("REVISION", "2026-07-28"), ("TRANSPORT", "stdio + in-memory"),
        ("SCHEMA", "JSON Schema 2020-12"), ("TELEMETRY", "W3C trace context"),
    ]):
        y = 2.32 + index * 0.46
        text(slide, key, 7.65, y, 1.5, 0.16, size=8, value_color="8FB6D6", bold=True)
        text(slide, value, 9.2, y - 0.02, 3.2, 0.22, size=11, value_color=WHITE, font="Cascadia Code")
    rect(slide, 7.35, 5.0, 5.35, 1.5, PAPER := "FAF9F8", line=LINE)
    code_line(slide, "async with Client(mcp) as client:", 7.6, 5.2, 5.0, size=11, value_color=MSBLUE_DK)
    code_line(slide, "    result = await client.call_tool(", 7.6, 5.5, 5.0, size=11, value_color=INK)
    code_line(slide, "        tool_name, arguments)", 7.6, 5.78, 5.0, size=11, value_color=INK)
    code_line(slide, "# bridge rejects non read-only tools", 7.6, 6.12, 5.0, size=10, value_color=GREEN)
    add_notes(slide, "Protocol revision 2026-07-28 is asserted in tests. The in-memory Client transport removed a redundant nested process and cut cold-start from ~19.7s to ~4.5s in-container.")


def d1_tools(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Tool contract & annotations", "Annotations are hints; policy is enforcement", 6)
    tools = [
        ("get_recall_notice", "READ", GREEN),
        ("locate_inventory", "READ", GREEN),
        ("get_supplier_status", "READ", GREEN),
        ("request_approval", "CONTROLLED WRITE", AMBER),
        ("quarantine_batch", "DESTRUCTIVE / IDEMPOTENT", RED),
        ("get_audit_events", "READ / SANITIZED", GREEN),
        ("reset_demo", "DEMO CONTROL", MUTED),
    ]
    for index, (name, mode, accent) in enumerate(tools):
        column = 0 if index < 4 else 1
        row = index if index < 4 else index - 4
        x = 0.58 + column * 4.3
        y = 1.72 + row * 1.05
        rect(slide, x, y, 4.0, 0.82, WHITE, line=LINE)
        rect(slide, x, y, 0.08, 0.82, accent)
        text(slide, name, x + 0.2, y + 0.13, 3.5, 0.22, size=12, font="Cascadia Code", bold=True)
        text(slide, mode, x + 0.2, y + 0.44, 3.6, 0.16, size=8, value_color=accent, bold=True)

    rect(slide, 9.3, 1.72, 3.4, 3.4, DEEP)
    text(slide, "ENFORCEMENT PATH", 9.55, 2.0, 3.0, 0.22, size=9, value_color=CYAN, bold=True)
    for index, item in enumerate([
        "Read tools reach agents",
        "Mutations require a token",
        "Token issued only by policy",
        "Bridge blocks write tools",
        "Audit hashes the approval",
    ]):
        y = 2.42 + index * 0.5
        rect(slide, 9.55, y + 0.03, 0.12, 0.12, MSBLUE)
        text(slide, item, 9.8, y - 0.05, 2.75, 0.4, size=11.5, value_color="E4ECF5")
    label(slide, "Demo 1", 0.58, 6.3, 0.8, RED)
    text(slide, "Discover schemas → fail safely → approve → quarantine → replay",
         1.5, 6.31, 11.0, 0.27, size=15, font="Segoe UI Semibold", bold=True)
    add_notes(slide, "An annotation such as destructive or read-only helps trusted clients render behavior. It is NOT authorization. The read-only bridge and the policy token are the actual enforcement, and both are covered by tests.")


def d1_agents(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Multi-agent pattern", "Specialize by evidence, not by personality", 7)
    agents = [
        ("01", "Recall Triage", "get_recall_notice", "Validate batch + risk tier"),
        ("02", "Inventory Impact", "locate_inventory", "Compute exposure from positions"),
        ("03", "Supplier / Compliance", "get_supplier_status", "Frame operational constraints"),
        ("04", "Supervisor", "NO TOOLS", "Synthesize the decision brief"),
    ]
    for index, (num, name, tool_name, outcome) in enumerate(agents):
        x = 0.56 + index * 3.12
        rect(slide, x, 1.86, 2.78, 3.15, WHITE, line=LINE)
        rect(slide, x, 1.86, 2.78, 0.08, MSBLUE if index < 3 else AMBER)
        text(slide, num, x + 0.2, 2.13, 0.42, 0.3, size=15, value_color=MSBLUE, font="Segoe UI Semibold", bold=True)
        text(slide, name, x + 0.2, 2.62, 2.35, 0.7, size=20, font="Segoe UI Semibold", bold=True)
        label(slide, tool_name, x + 0.2, 3.5, 2.3, GREEN_SOFT if index < 3 else AMBER_SOFT, GREEN if index < 3 else AMBER)
        text(slide, outcome, x + 0.2, 4.1, 2.3, 0.7, size=12.5, value_color=MUTED)
        if index < 3:
            text(slide, "›", x + 2.82, 3.15, 0.26, 0.3, size=22, value_color=LINE, font="Segoe UI Semibold", bold=True)
    rect(slide, 0.56, 5.42, 12.1, 0.95, DEEP)
    text(slide, "LEAST PRIVILEGE", 0.86, 5.66, 2.0, 0.2, size=9, value_color=CYAN, bold=True)
    text(slide, "No agent receives request_approval or quarantine_batch.", 2.7, 5.56, 6.0, 0.34, size=16, value_color=WHITE, font="Segoe UI Semibold", bold=True)
    text(slide, "The application keeps the credential.", 9.0, 5.6, 3.4, 0.27, size=12.5, value_color="C7D6E6")
    add_notes(slide, "Sequential because each step depends on validated evidence from the previous step — not agents for the sake of agents. The Supervisor is tool-free and receives full workflow context.")


def d1_safety(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Determinism & safety", "Make the consequential path boring and provable", 8)
    cards = [
        ("APPROVAL GATE", "Quarantine requires a named approver and a policy-issued token. Prompts cannot mint authority.", MSBLUE),
        ("IDEMPOTENT MUTATION", "Quarantine is batch-scoped and keyed; a retry changes zero additional positions.", GREEN),
        ("SANITIZED ERRORS", "Tool failures return a generic client message; stderr is discarded from responses.", AMBER),
        ("AUDIT WITHOUT SECRETS", "Approval appears as a hash only; no usable credential is ever emitted.", RED),
    ]
    for index, (title_value, body, accent) in enumerate(cards):
        x = 0.58 + (index % 2) * 6.14
        y = 1.75 + (index // 2) * 2.15
        rect(slide, x, y, 5.9, 1.9, WHITE, line=LINE)
        rect(slide, x, y, 0.09, 1.9, accent)
        text(slide, title_value, x + 0.28, y + 0.22, 5.4, 0.3, size=15, font="Segoe UI Semibold", bold=True)
        text(slide, body, x + 0.28, y + 0.72, 5.4, 1.0, size=13, value_color=MUTED)
    add_notes(slide, "These four properties are what turn a plausible prototype into something you can operate. Each is covered by a focused test in the suite.")


def d1_demo(prs):
    demo_slide(
        prs,
        number=9,
        demo_name="Demo 1 · Local MCP + multi-agent control",
        goal="Prove the workflow can make progress, fail clearly, recover safely, and complete a consequential mutation without letting the model manufacture authority.",
        steps=[
            "Start the Control Tower; show ACTIVE RECALL, 4 locations, 2,196 units.",
            "Run analysis; follow Triage → Inventory → Supplier → Supervisor.",
            "Attempt quarantine before approval; show the deterministic denial.",
            "Approve as 'Asha Rao, Responsible Pharmacist'; quarantine the batch.",
            "Replay quarantine; then open the sanitized audit trail.",
        ],
        show=[
            "MCP tool calls per agent (read-only only)",
            "First call: 4 positions / 2,196 units",
            "Replay: 0 additional positions changed",
            "Audit: hashed approval ID, no credential",
            "Terminal: pytest — all local tests passing",
        ],
        fallback="assets/control-tower-quarantined.png + narrate the four proof points.",
        command="python -m uvicorn control_tower_api:app --app-dir ./src --port 8091",
    )


def d1_proof(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Demo 1 proof", "The same mutation can be retried without being repeated", 10)
    add_picture(slide, "control-tower-quarantined.png", 0.58, 1.66, 8.0, 4.96)
    rect(slide, 8.88, 1.66, 3.82, 4.96, DEEP)
    label(slide, "Local result", 9.2, 2.0, 1.3, GREEN)
    proof = [
        ("01", "Approval named", "Asha Rao, Responsible Pharmacist"),
        ("02", "First call", "4 positions / 2,196 units"),
        ("03", "Replay", "0 additional positions"),
        ("04", "Audit", "Hashed approval ID only"),
    ]
    for index, (num, title_value, detail) in enumerate(proof):
        y = 2.56 + index * 0.86
        text(slide, num, 9.2, y, 0.42, 0.22, size=11, value_color=CYAN, bold=True)
        text(slide, title_value, 9.72, y - 0.03, 2.7, 0.25, size=14, value_color=WHITE, bold=True)
        text(slide, detail, 9.72, y + 0.27, 2.75, 0.3, size=10, value_color="A8C0D6")
    text(slide, "16 local tests passing", 9.2, 6.08, 3.0, 0.25, size=13, value_color=CYAN, font="Cascadia Code", bold=True)
    add_notes(slide, "Summarize only the proof: approval named, first mutation changed four positions, replay changed zero, audit output contained no usable credential. Then transition: production adds identity, routing, tracing, evaluation, and deployment governance — Deck 2.")


def d1_close(prs):
    closing_slide(
        prs,
        title_value="Build the system\naround the model",
        rules=[
            ("01", "Put authority below prompts"),
            ("02", "Give each agent the minimum tool surface"),
            ("03", "Design every mutation for retries"),
            ("04", "Treat protocol and tests as product code"),
        ],
        proof_headline="2,196",
        proof_caption="units quarantined once",
        proof_items=["MCP 2026-07-28", "4 least-privilege agents", "Named human approval", "Idempotent replay", "16 tests green"],
        resources="learn.microsoft.com/azure/ai-foundry · modelcontextprotocol.io · Reproduce locally: python -m pytest",
    )


# ---------------------------------------------------------------------------
# Deck 2 — Demo 2: Microsoft Foundry Hosted Agent production
# ---------------------------------------------------------------------------

def d2_title(prs):
    title_slide(
        prs,
        demo_tag="Demo 2 · Managed production",
        title_value="Operating the Agent",
        subtitle="The same modular system as a Microsoft Foundry Hosted Agent — reproducible, hardened, observable, and governed",
        chips=["Hosted Agent", "Model Router", "Managed identity", "azd lifecycle"],
        disclaimer=DISCLAIMER,
    )


def d2_bridge(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "From correct to operable", "What production adds on top of a working prototype", 2)
    rows = [
        ("Runtime", "Local .venv process", "Foundry Hosted Agent container"),
        ("Identity", "Developer credentials", "Instance managed identity"),
        ("Model", "Direct client call", "Model Router deployment contract"),
        ("Build", "pip install", "Hash-locked, non-root, CVE-scanned image"),
        ("Evidence", "Local logs", "App Insights trace correlation + eval"),
        ("Lifecycle", "Run script", "azd provision → deploy → invoke"),
    ]
    text(slide, "PROTOTYPE", 3.5, 1.6, 3.0, 0.22, size=9, value_color=MUTED, bold=True)
    text(slide, "MANAGED PRODUCTION", 8.0, 1.6, 4.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, (dim, before, after) in enumerate(rows):
        y = 2.0 + index * 0.78
        rect(slide, 0.58, y, 2.7, 0.62, DEEP)
        text(slide, dim, 0.78, y + 0.16, 2.4, 0.3, size=13, value_color=WHITE, font="Segoe UI Semibold", bold=True)
        rect(slide, 3.4, y, 4.3, 0.62, WHITE, line=LINE)
        text(slide, before, 3.6, y + 0.17, 4.0, 0.3, size=12, value_color=MUTED)
        rect(slide, 7.85, y, 4.85, 0.62, WHITE, line=LINE)
        rect(slide, 7.85, y, 0.08, 0.62, MSBLUE)
        text(slide, after, 8.05, y + 0.17, 4.5, 0.3, size=12, value_color=INK, bold=True)
    add_notes(slide, "Nothing about the agent logic changes. What changes is the runtime, identity, model contract, supply chain, evidence, and lifecycle — the operability surface.")


def d2_yaml(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "azure.yaml is the source of truth", "Deployment intent lives in Git, not in a console", 3)
    rect(slide, 0.58, 1.7, 6.5, 4.9, DEEP)
    lines = [
        ("services:", WHITE),
        ("  ai-project:", CYAN),
        ("    host: azure.ai.project", "C7D6E6"),
        ("    deployments:", "C7D6E6"),
        ("      - name: caldova-model-router", "C7D6E6"),
        ("        model: { name: model-router,", "9FB8D0"),
        ("          version: 2025-11-18 }", "9FB8D0"),
        ("  agent-framework-workflows-responses:", CYAN),
        ("    host: azure.ai.agent", "C7D6E6"),
        ("    language: docker", "C7D6E6"),
        ("    docker: { path: Dockerfile,", "9FB8D0"),
        ("      remoteBuild: true }", "9FB8D0"),
        ("    uses: [ai-project]", "C7D6E6"),
        ("    container: { resources:", "9FB8D0"),
        ("      { cpu: '0.5', memory: 1Gi } }", "9FB8D0"),
        ("infra: { provider: microsoft.foundry }", GREEN),
    ]
    for index, (value, col) in enumerate(lines):
        code_line(slide, value, 0.8, 1.95 + index * 0.28, 6.1, size=10.5, value_color=col)

    text(slide, "WHY IT MATTERS", 7.35, 1.75, 5.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, item in enumerate([
        "One manifest owns model, agent and infra",
        "remoteBuild → ACR builds the Dockerfile",
        "No codeConfiguration/image → azd owns build",
        "container.resources pins 0.5 CPU / 1 GiB",
        "Reviewable, diffable, promotable in Git",
    ]):
        bullet_row(slide, item, 7.35, 2.12 + index * 0.6, 5.3, size=13, accent=MSBLUE)
    rect(slide, 7.35, 5.3, 5.35, 1.2, GREEN_SOFT)
    text(slide, "azd provision reads infra; azd deploy reads the agent block and registers an immutable version.",
         7.55, 5.5, 5.0, 0.9, size=12.5, value_color=INK)
    add_notes(slide, "The Foundry provider marshals the deployments into the template. We hit and fixed a real escaping defect where the provider passed JSON with escaped quotes — the template now defensively un-escapes before json().")


def d2_supplychain(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Container & supply chain", "The image is a security boundary, so we hardened it", 4)
    cards = [
        ("DUAL INTERPRETER", "System Python runs the Hosted Agent; an isolated /opt/caldova-mcp runs MCP 2.1.0. Two runtimes, one image.", MSBLUE),
        ("NON-ROOT", "Runs as UID 10001 with least file access — verified id -u = 10001 inside the container.", GREEN),
        ("HASH-LOCKED DEPS", "requirements.lock + requirements-mcp.lock pinned with hashes, generated in the Linux base image.", PURPLE),
        ("CVE-CLEAN + SBOM", "Build tooling removed; Trivy reports zero fixable HIGH/CRITICAL; SBOM retained as evidence.", RED),
    ]
    for index, (title_value, body, accent) in enumerate(cards):
        x = 0.58 + (index % 2) * 6.14
        y = 1.75 + (index // 2) * 2.15
        rect(slide, x, y, 5.9, 1.9, WHITE, line=LINE)
        rect(slide, x, y, 0.09, 1.9, accent)
        text(slide, title_value, x + 0.28, y + 0.22, 5.4, 0.3, size=15, font="Segoe UI Semibold", bold=True)
        text(slide, body, x + 0.28, y + 0.72, 5.4, 1.0, size=12.5, value_color=MUTED)
    add_notes(slide, "Cold start under 0.5 CPU / 1 GiB measured ~4.5–6.6s in-container; a bounded, env-configurable timeout and a concurrency semaphore protect the constrained runtime.")


def d2_router(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Model Router", "One deployment contract, governed runtime choices", 5)
    rect(slide, 0.58, 1.75, 4.25, 4.85, DEEP)
    text(slide, "CALDOVA MODEL ROUTER", 0.9, 2.08, 3.6, 0.25, size=10, value_color=CYAN, bold=True)
    text(slide, "caldova-model-router", 0.9, 2.5, 3.4, 0.43, size=24, value_color=WHITE, font="Segoe UI Semibold", bold=True)
    text(slide, "model-router / 2025-11-18", 0.9, 3.02, 3.2, 0.25, size=12, value_color="A8C0D6", font="Cascadia Code")
    for index, (key, value) in enumerate([("PROFILE", "Balanced"), ("SKU", "GlobalStandard"), ("CAPACITY", "10"), ("REGION", "northcentralus")]):
        y = 3.68 + index * 0.5
        text(slide, key, 0.9, y, 1.5, 0.18, size=8, value_color="7F9DB5", bold=True)
        text(slide, value, 2.4, y - 0.02, 2.0, 0.22, size=11, value_color=WHITE, font="Cascadia Code")

    text(slide, "ONE CONTRACT, MANY MODELS", 5.1, 1.75, 6.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, item in enumerate([
        "Every agent reads AZURE_AI_MODEL_DEPLOYMENT_NAME — no hardcoded model",
        "Router selects an eligible model per request under one deployment",
        "Switch or scale models without touching agent code",
        "Telemetry surfaces the routed model where the platform returns it",
    ]):
        bullet_row(slide, item, 5.1, 2.15 + index * 0.72, 7.5, size=13, accent=MSBLUE, height=0.65)
    rect(slide, 5.1, 5.2, 7.6, 1.35, GREEN_SOFT)
    text(slide, "Evaluate quality, latency, routed-model distribution and cost on representative prompts before making production claims.",
         5.32, 5.42, 7.2, 1.0, size=13, value_color=INK)
    add_notes(slide, "Model Router keeps a single deployment contract while the platform picks an eligible model. This is how you decouple the workflow from any one model's lifecycle.")


def d2_rbac(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Identity & RBAC", "Control plane and data plane are different jobs", 6)
    rect(slide, 0.58, 1.75, 5.75, 4.5, WHITE, line=LINE)
    label(slide, "Control plane", 0.9, 2.05, 1.6, INK)
    text(slide, "Deployment identity", 0.9, 2.5, 4.0, 0.4, size=22, font="Segoe UI Semibold", bold=True)
    for index, item in enumerate(["Provision Foundry + supporting resources", "Build/push image to ACR (scoped)", "Register immutable agent version", "Production CI target: workload federation"]):
        bullet_row(slide, item, 0.9, 3.2 + index * 0.6, 5.1, size=13, accent=MSBLUE)

    rect(slide, 6.8, 1.75, 5.9, 4.5, DEEP)
    label(slide, "Data plane", 7.13, 2.05, 1.4, MSBLUE)
    text(slide, "Hosted Agent identity", 7.13, 2.5, 4.5, 0.4, size=22, value_color=WHITE, font="Segoe UI Semibold", bold=True)
    for index, item in enumerate(["Invoke Model Router", "Call only required operational services", "Read/write only required data", "Emit telemetry without credentials"]):
        rect(slide, 7.15, 3.24 + index * 0.6, 0.12, 0.12, CYAN)
        text(slide, item, 7.42, 3.16 + index * 0.6, 5.1, 0.4, size=13, value_color="E4ECF5")
    rect(slide, 0.58, 6.42, 12.12, 0.4, RED_SOFT)
    text(slide, "The presenter identity is never the production runtime identity.", 0.8, 6.49, 11.6, 0.2, size=12, value_color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Deployment permissions can change infrastructure; runtime permissions can invoke and access data. They must be independently testable and scoped. Verified data-plane principal on the deployed agent.")


def d2_observability(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Observability & evaluation", "One trace answers the whole question", 7)
    stages = [
        ("UI", "Control Tower request", MSBLUE),
        ("WORKFLOW", "Which specialist step", MSBLUE),
        ("MODEL", "Routed model + latency", GREEN),
        ("MCP", "Which tool, what result", GREEN),
        ("POLICY", "Deterministic outcome", RED),
    ]
    for index, (tag, detail, accent) in enumerate(stages):
        x = 0.58 + index * 2.45
        rect(slide, x, 1.9, 2.2, 2.2, WHITE, line=LINE)
        rect(slide, x, 1.9, 2.2, 0.09, accent)
        text(slide, f"0{index + 1}", x + 0.2, 2.2, 0.5, 0.25, size=13, value_color=accent, font="Segoe UI Semibold", bold=True)
        text(slide, tag, x + 0.2, 2.68, 1.9, 0.3, size=15, font="Segoe UI Semibold", bold=True)
        text(slide, detail, x + 0.2, 3.12, 1.9, 0.8, size=11, value_color=MUTED)
        if index < 4:
            text(slide, "›", x + 2.24, 2.7, 0.24, 0.3, size=20, value_color=LINE, font="Segoe UI Semibold", bold=True)
    rect(slide, 0.58, 4.5, 12.12, 1.0, DEEP)
    text(slide, "CORRELATION IS THE DEBUGGING INTERFACE", 0.85, 4.72, 5.0, 0.25, size=10, value_color=CYAN, bold=True)
    text(slide, "One App Insights trace ties the UI request to the routed model, the MCP tool call, and the policy outcome.",
         0.85, 5.05, 11.4, 0.35, size=15, value_color=WHITE, font="Segoe UI Semibold", bold=True)
    text(slide, "Before production claims: measure response quality, latency, routed-model distribution, and estimated cost — azd ai agent eval generate.",
         0.58, 5.75, 12.1, 0.6, size=12.5, value_color=INK)
    add_notes(slide, "A model response alone is not enough. Correlate the UI request, workflow step, routed model, MCP call, and deterministic policy outcome in a single trace.")


def d2_lifecycle(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Reproducible lifecycle", "Anyone can stand this up from a clean clone", 8)
    steps = [
        ("01", "BOOTSTRAP", "python scripts/setup_azd_env.py", "Creates the azd env, detects the principal, sets non-secret values"),
        ("02", "PROVISION", "azd provision", "Foundry project, Model Router, ACR, App Insights, Log Analytics"),
        ("03", "DEPLOY", "azd deploy", "ACR remote build → immutable Hosted Agent version"),
        ("04", "INVOKE", "azd ai agent invoke …", "Batch B-2408-AX7 → 2,196 units + correlated trace"),
    ]
    for index, (num, tag, cmd, detail) in enumerate(steps):
        y = 1.8 + index * 1.12
        rect(slide, 0.58, y, 12.12, 0.95, WHITE, line=LINE)
        rect(slide, 0.58, y, 0.08, 0.95, MSBLUE)
        text(slide, num, 0.8, y + 0.3, 0.6, 0.3, size=16, value_color=MSBLUE, font="Segoe UI Semibold", bold=True)
        text(slide, tag, 1.5, y + 0.16, 2.4, 0.3, size=14, font="Segoe UI Semibold", bold=True)
        rect(slide, 1.5, y + 0.52, 3.4, 0.32, "FAF9F8", line=LINE)
        code_line(slide, cmd, 1.62, y + 0.58, 3.2, size=10, value_color=MSBLUE_DK)
        text(slide, detail, 5.1, y + 0.32, 7.4, 0.5, size=12.5, value_color=MUTED)
    add_notes(slide, "No .azure or identity state is committed. A collaborator runs the bootstrap and gets an equivalent environment. All high-severity blockers were cleared before this ran.")


def d2_demo(prs):
    demo_slide(
        prs,
        number=9,
        demo_name="Demo 2 · Foundry Hosted Agent in production",
        goal="Operate the same modular system securely and reproducibly at scale — managed runtime, managed identity, routed model, and a correlated trace.",
        steps=[
            "azd ai agent show — confirm version 1 is active.",
            "Open the Foundry playground for the deployed agent.",
            "Invoke batch B-2408-AX7 with the recall prompt.",
            "Open the correlated App Insights trace.",
            "Show the data-plane managed identity and scoped role.",
        ],
        show=[
            "Agent status active; immutable version 1",
            "Response: same 4 locations / 2,196 units",
            "Routed model + latency in the trace",
            "MCP tool activity inside the trace",
            "Managed identity ≠ presenter identity",
        ],
        fallback="Narrate azure.yaml as the deployment contract; do not claim live routing without the trace on screen.",
        command="azd ai agent invoke agent-framework-workflows-responses \"…B-2408-AX7…\"",
    )


def d2_proof(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Demo 2 proof", "Deployed, invoked, and correlated", 10)
    rect(slide, 0.58, 1.72, 6.0, 4.9, DEEP)
    text(slide, "DEPLOYED AGENT", 0.85, 2.0, 4.0, 0.22, size=9, value_color=CYAN, bold=True)
    for index, (key, value) in enumerate([
        ("Version", "agent-framework-workflows-responses:1"),
        ("Status", "active"),
        ("Identity", "instance managed identity"),
        ("Registry", "crymwyto…azurecr.io (scoped)"),
        ("Endpoint", "…/protocols/openai/responses"),
    ]):
        y = 2.4 + index * 0.62
        text(slide, key, 0.85, y, 1.6, 0.2, size=9, value_color="7F9DB5", bold=True)
        text(slide, value, 0.85, y + 0.22, 5.5, 0.3, size=12, value_color=WHITE, font="Cascadia Code")

    rect(slide, 6.75, 1.72, 5.95, 4.9, WHITE, line=LINE)
    text(slide, "VERIFIED INVOCATION", 7.0, 2.0, 4.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, item in enumerate([
        "Batch B-2408-AX7 → 2,196 units, 4 locations",
        "Per-location totals match the local run",
        "Supplier acknowledgement reported",
        "Refused quarantine without named approval",
        "Trace captured; first byte ~5s",
    ]):
        bullet_row(slide, item, 7.0, 2.45 + index * 0.72, 5.4, size=13, accent=GREEN, height=0.65)
    add_notes(slide, "One honest caveat: in the live run the supervisor summary under-represented the supplier/compliance step — a routing/prompt nuance to tighten, not a code defect (local tests pass). Investigate via the trace before the session.")


def d2_close(prs):
    closing_slide(
        prs,
        title_value="Managed runtime,\nengineered controls",
        rules=[
            ("01", "Keep the manifest the source of truth"),
            ("02", "Separate control-plane and data-plane identity"),
            ("03", "Harden the image; lock the supply chain"),
            ("04", "Correlate one trace across every layer"),
        ],
        proof_headline="v1",
        proof_caption="active hosted agent",
        proof_items=["Reproducible azd lifecycle", "Model Router contract", "Managed identity runtime", "CVE-clean image + SBOM", "Correlated trace evidence"],
        resources="learn.microsoft.com/azure/ai-foundry/agents · aka.ms/azd · Foundry playground linked in azd output",
    )


DECK1 = [
    d1_title, d1_problem, d1_scenario, d1_architecture, d1_mcp, d1_tools,
    d1_agents, d1_safety, d1_demo, d1_proof, d1_close,
]
DECK2 = [
    d2_title, d2_bridge, d2_yaml, d2_supplychain, d2_router, d2_rbac,
    d2_observability, d2_lifecycle, d2_demo, d2_proof, d2_close,
]


def build_one(builders, *, title, subject, output_name):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = title
    prs.core_properties.subject = subject
    prs.core_properties.author = "Caldova Recall Control Tower"
    for builder in builders:
        builder(prs)
    output = ROOT / output_name
    prs.save(output)
    return output


def build():
    return [
        build_one(
            DECK1,
            title="Engineering the Agent — Caldova Demo 1 (MCP + Multi-Agent)",
            subject="MCP v2 protocol contract, multi-agent least privilege, deterministic authority",
            output_name="caldova-demo1-mcp-multiagent.pptx",
        ),
        build_one(
            DECK2,
            title="Operating the Agent — Caldova Demo 2 (Microsoft Foundry)",
            subject="Microsoft Foundry Hosted Agent, Model Router, managed identity, observability",
            output_name="caldova-demo2-foundry-production.pptx",
        ),
    ]


if __name__ == "__main__":
    for path in build():
        print(f"wrote {path.name}")
