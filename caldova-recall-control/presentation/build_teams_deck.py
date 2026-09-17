"""Build the Caldova "Deploy as a Microsoft Teams hosted agent" deck.

Covers deployment and implementation of the recall solution as a Microsoft Teams
app backed by the Foundry Hosted Agent, plus a demo scenario for deploying it.
Aimed at engineering + IT / Teams administrators. Run: `python build_teams_deck.py`.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from deck_common import (
    AMBER, AMBER_SOFT, BLUE_SOFT, CYAN, DEEP, GREEN, GREEN_SOFT, INK, LINE,
    MSBLUE, MSBLUE_DK, MUTED, PURPLE, RED, RED_SOFT, SLIDE_H, SLIDE_W, WHITE,
    add_notes, bullet_row, closing_slide, code_line, demo_slide, heading,
    label, line, node, rect, text, title_slide,
)


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
DISCLAIMER = "Fictional pharmaceutical recall. Synthetic operational data. Batch B-2408-AX7."


def add_picture(slide, filename, x, y, w, h):
    return slide.shapes.add_picture(str(ASSETS / filename), Inches(x), Inches(y), Inches(w), Inches(h))


def s_title(prs):
    title_slide(
        prs,
        demo_tag="Teams deployment · Engineering + IT",
        title_value="Caldova in Microsoft Teams",
        subtitle="Deliver the governed recall workflow inside Teams — a hosted web tab backed by the Foundry Hosted Agent, published and controlled from the Teams admin center",
        chips=["Teams tab + SSO", "Foundry Hosted Agent", "Entra ID", "Admin-center governed"],
        disclaimer=DISCLAIMER,
    )


def s_goal(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Why Teams", "Meet operators where they already work", 2)
    text(slide, "THE SAME GOVERNED WORKFLOW, IN TEAMS", 0.58, 1.7, 7.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, item in enumerate([
        "Operators run recall triage, approval, and quarantine inside a Teams tab",
        "No new client to install — the app is pinned by policy for the target group",
        "Single sign-on with the user's Entra identity; no separate credentials",
        "Authority stays deterministic and below the model, exactly as in the app",
        "Adoption, usage, and audit are visible to IT through standard tooling",
    ]):
        bullet_row(slide, item, 0.58, 2.08 + index * 0.6, 7.2, size=13, accent=MSBLUE)

    rect(slide, 8.2, 1.7, 4.5, 4.6, DEEP)
    text(slide, "TWO AUDIENCES, ONE PACKAGE", 8.45, 2.0, 4.0, 0.22, size=9, value_color=CYAN, bold=True)
    text(slide, "Engineering", 8.45, 2.42, 4.0, 0.3, size=15, value_color=WHITE, font="Segoe UI Semibold", bold=True)
    text(slide, "Builds the hosted agent, web tab, and SSO wiring.", 8.45, 2.78, 4.0, 0.5, size=12, value_color="C7D6E6")
    text(slide, "IT / Teams admin", 8.45, 3.5, 4.0, 0.3, size=15, value_color=WHITE, font="Segoe UI Semibold", bold=True)
    text(slide, "Publishes, targets, governs, and operates the app for the org.", 8.45, 3.86, 4.0, 0.6, size=12, value_color="C7D6E6")
    rect(slide, 8.45, 4.9, 4.0, 1.15, "122A47")
    text(slide, "This deck hands off cleanly from the build team to the Teams administrator.", 8.65, 5.08, 3.7, 0.9, size=12, value_color="E4ECF5")
    add_notes(slide, "The point of Teams delivery is reach and governance: the same deterministic recall controls, surfaced where operators already are, under IT policy.")


def s_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Architecture in Teams", "Client to tab to hosted agent, with identity at every hop", 3)
    node(slide, "Teams client", "Desktop / web / mobile", 0.55, 2.5, 1.9, 1.0, MSBLUE)
    node(slide, "Teams tab (SSO)", "Entra token for the user", 2.62, 2.5, 2.0, 1.0, MSBLUE)
    node(slide, "Hosted web app", "App Service / Container Apps", 4.8, 2.5, 2.1, 1.0, GREEN)
    node(slide, "Foundry Hosted Agent", "Responses protocol", 7.1, 1.6, 2.2, 1.0, GREEN)
    node(slide, "Model Router", "caldova-model-router", 7.1, 3.3, 2.2, 1.0, GREEN)
    node(slide, "MCP + policy", "Approval / idempotency / audit", 9.55, 2.5, 2.15, 1.0, RED)
    for x1, y1, x2, y2 in [
        (2.45, 3.0, 2.62, 3.0), (4.62, 3.0, 4.8, 3.0), (6.9, 3.0, 7.1, 2.1),
        (6.9, 3.0, 7.1, 3.8), (9.3, 2.1, 9.55, 3.0), (9.3, 3.8, 9.55, 3.2),
    ]:
        line(slide, x1, y1, x2, y2, MSBLUE if y2 < 3.9 else RED, 1.5, dash=False)
    text(slide, "Entra SSO", 2.7, 2.15, 1.8, 0.2, size=9, value_color=MSBLUE, bold=True)
    text(slide, "managed identity", 4.85, 2.15, 2.0, 0.2, size=9, value_color=GREEN, bold=True)
    text(slide, "mutate only after approval", 9.5, 3.75, 2.2, 0.2, size=9, value_color=RED, bold=True)
    rect(slide, 0.55, 6.15, 12.15, 0.55, BLUE_SOFT)
    text(slide, "The Teams tab passes the user's Entra identity to a private hosted app; the agent runs on its own managed identity.",
         0.78, 6.29, 11.7, 0.22, size=14, value_color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_notes(slide, "Two identities: the USER (via Teams SSO to the web app) and the AGENT (managed identity to Model Router and services). Neither is the other.")


def s_patterns(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Integration pattern", "Tab-embedded web app vs bot — and what we ship", 4)
    cards = [
        ("TAB (CHOSEN)", "Embed the existing Control Tower web UI as a Teams tab. Full fidelity UI, approval dialogs, and audit view with minimal new code.", MSBLUE),
        ("BOT / MESSAGE EXT", "Conversational access to the agent in chat. Useful later for quick queries, but weaker for consequential approval UX.", MUTED),
    ]
    for index, (title_value, body, accent) in enumerate(cards):
        x = 0.58 + index * 6.15
        rect(slide, x, 1.75, 5.9, 1.9, WHITE, line=LINE)
        rect(slide, x, 1.75, 0.09, 1.9, accent)
        text(slide, title_value, x + 0.28, 1.98, 5.4, 0.3, size=15, font="Segoe UI Semibold", bold=True)
        text(slide, body, x + 0.28, 2.48, 5.4, 1.0, size=12.5, value_color=MUTED)
    text(slide, "WHY A TAB FOR THIS WORKLOAD", 0.58, 4.0, 7.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, item in enumerate([
        "Consequential approval needs an explicit, auditable UI — not a chat turn",
        "Reuses the hardened web app and its authorization, unchanged",
        "SSO via webApplicationInfo gives seamless, tenant-only access",
        "A bot/custom engine agent can be added later against the same hosted agent",
    ]):
        bullet_row(slide, item, 0.58, 4.36 + index * 0.56, 12.0, size=13, accent=MSBLUE)
    add_notes(slide, "We surface the web app as a tab because approval is a governed, evidence-producing action. The bot pattern is a future add-on, not the primary control surface.")


def s_identity(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Identity & SSO", "Entra ID app registration and Teams single sign-on", 5)
    rect(slide, 0.58, 1.7, 6.4, 4.6, DEEP)
    text(slide, "ENTRA APP REGISTRATION", 0.85, 2.0, 5.0, 0.22, size=9, value_color=CYAN, bold=True)
    for index, item in enumerate([
        "Single-tenant app for the web tab",
        "Application ID URI: api://<domain>/<app-id>",
        "Expose scope: access_as_user",
        "Pre-authorize Teams client app IDs",
        "Admin consent: Graph User.Read",
        "Require user assignment; assign target group",
    ]):
        rect(slide, 0.85, 2.45 + index * 0.55, 0.12, 0.12, MSBLUE)
        text(slide, item, 1.1, 2.37 + index * 0.55, 5.7, 0.4, size=12.5, value_color="E4ECF5")

    text(slide, "MANIFEST SSO BINDING", 7.35, 1.7, 5.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    rect(slide, 7.35, 2.02, 5.35, 1.85, "FAF9F8", line=LINE)
    code_line(slide, "\"webApplicationInfo\": {", 7.55, 2.2, 5.0, size=11, value_color=MSBLUE_DK)
    code_line(slide, "  \"id\": \"<entra-app-client-id>\",", 7.55, 2.5, 5.0, size=11, value_color=INK)
    code_line(slide, "  \"resource\": \"api://<domain>/<id>\"", 7.55, 2.8, 5.0, size=11, value_color=INK)
    code_line(slide, "}", 7.55, 3.1, 5.0, size=11, value_color=MSBLUE_DK)
    text(slide, "CONTROL VS DATA PLANE", 7.35, 4.1, 5.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    text(slide, "User identity (Teams SSO) authorizes the web tab. The agent uses its own managed identity for Model Router and services. Apply Conditional Access to the enterprise app.",
         7.35, 4.42, 5.35, 1.6, size=12.5, value_color=INK)
    add_notes(slide, "SSO is the crux: pre-authorize the Teams client IDs for the access_as_user scope and bind webApplicationInfo so the tab exchanges the Teams token for the app silently.")


def s_hosting(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Host it securely", "The tab is only as safe as the app behind it", 6)
    cards = [
        ("ENTRA-ONLY ACCESS", "App Service Easy Auth (or equivalent) bound to the Entra app; deny anonymous access; HTTPS + custom domain + TLS.", MSBLUE),
        ("AUTHORIZE MUTATIONS", "The demo's approval/quarantine/reset endpoints have no auth and bind to localhost. Add Entra role/group checks before any org exposure.", RED),
        ("NETWORK CONTROLS", "Private endpoints / IP allow-list / WAF per policy. No public anonymous surface for the app or the agent.", PURPLE),
        ("OBSERVABILITY", "Wire the web app and hosted agent to enterprise Application Insights / Log Analytics; keep secrets out of logs.", GREEN),
    ]
    for index, (title_value, body, accent) in enumerate(cards):
        x = 0.58 + (index % 2) * 6.15
        y = 1.75 + (index // 2) * 2.15
        rect(slide, x, y, 5.9, 1.9, WHITE, line=LINE)
        rect(slide, x, y, 0.09, 1.9, accent)
        text(slide, title_value, x + 0.28, y + 0.22, 5.4, 0.3, size=14.5, font="Segoe UI Semibold", bold=True)
        text(slide, body, x + 0.28, y + 0.7, 5.4, 1.05, size=12, value_color=MUTED)
    add_notes(slide, "The single most important production gate: the mutation endpoints must gain authorization. Everything else is standard secure web hosting.")


def s_package(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Teams app package", "A manifest, two icons, one tab", 7)
    rect(slide, 0.58, 1.7, 6.5, 4.6, DEEP)
    for index, (value, col) in enumerate([
        ("{", WHITE),
        ("  \"manifestVersion\": \"1.17\",", CYAN),
        ("  \"id\": \"<new-app-guid>\",", "C7D6E6"),
        ("  \"name\": { \"short\": \"Caldova Recall\" },", "C7D6E6"),
        ("  \"staticTabs\": [{", "C7D6E6"),
        ("    \"entityId\": \"caldova-control-tower\",", "9FB8D0"),
        ("    \"contentUrl\": \"https://<host>/\",", "9FB8D0"),
        ("    \"scopes\": [\"personal\"] }],", "9FB8D0"),
        ("  \"webApplicationInfo\": {", CYAN),
        ("    \"id\": \"<entra-app-id>\",", "9FB8D0"),
        ("    \"resource\": \"api://<domain>/<id>\" },", "9FB8D0"),
        ("  \"validDomains\": [\"<host-domain>\"]", "C7D6E6"),
        ("}", WHITE),
    ]):
        code_line(slide, value, 0.8, 1.95 + index * 0.32, 6.1, size=11, value_color=col)

    text(slide, "PACKAGE CONTENTS", 7.35, 1.75, 5.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    for index, item in enumerate([
        "manifest.json (schema 1.17+)",
        "color icon 192x192 PNG",
        "outline icon 32x32 PNG",
        "Zip the three into the app package",
        "Validate in Teams Developer Portal",
    ]):
        bullet_row(slide, item, 7.35, 2.12 + index * 0.6, 5.3, size=13, accent=MSBLUE)
    rect(slide, 7.35, 5.25, 5.35, 1.0, AMBER_SOFT)
    text(slide, "Static tab = personal scope. Add a configurable tab for team/channel scope.",
         7.55, 5.42, 5.0, 0.75, size=12, value_color=INK)
    add_notes(slide, "The package is small: manifest + two icons. The contentUrl points at the hosted web app; webApplicationInfo enables SSO; validDomains must include the host.")


def s_publish(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Publish & govern", "Teams admin center owns availability and policy", 8)
    steps = [
        ("01", "UPLOAD", "Teams apps -> Manage apps -> Upload new app; set status Allowed"),
        ("02", "REVIEW", "Check Permissions and Data handling; record data flow"),
        ("03", "PERMIT", "App permission policy allows the app for the target group"),
        ("04", "PIN", "App setup policy pins / pre-installs for the group via assignment"),
        ("05", "GOVERN", "Purview labels, retention, audit logging; Conditional Access"),
    ]
    for index, (num, tag, detail) in enumerate(steps):
        y = 1.8 + index * 0.9
        rect(slide, 0.58, y, 12.12, 0.75, WHITE, line=LINE)
        rect(slide, 0.58, y, 0.08, 0.75, MSBLUE)
        text(slide, num, 0.8, y + 0.22, 0.6, 0.3, size=15, value_color=MSBLUE, font="Segoe UI Semibold", bold=True)
        text(slide, tag, 1.45, y + 0.22, 2.3, 0.3, size=14, font="Segoe UI Semibold", bold=True)
        text(slide, detail, 3.9, y + 0.24, 8.6, 0.4, size=12.5, value_color=MUTED)
    add_notes(slide, "Availability is policy-driven: allow the app, then target it to a security group with permission + setup policies. Governance (Purview, CA, audit) rides alongside.")


def s_demo(prs):
    demo_slide(
        prs,
        number=9,
        demo_name="Demo · Deploy the hosted agent to Teams",
        goal="Take the recall solution from a deployed Foundry Hosted Agent to a governed Teams tab a target group can use — end to end.",
        steps=[
            "azd provision && azd deploy — confirm agent version active.",
            "Publish the hosted web app; enable Entra Easy Auth + mutation authz.",
            "Register the Entra app; expose access_as_user; pre-auth Teams client IDs.",
            "Build the app package (manifest + icons); validate in Developer Portal.",
            "Upload in Teams admin center; assign permission + setup policy to the pilot group.",
        ],
        show=[
            "Agent status active; endpoint reachable",
            "Teams tab loads with silent SSO",
            "Run analysis -> approve -> quarantine in Teams",
            "Audit trail + App Insights trace correlate",
            "App pinned for the target group by policy",
        ],
        fallback="Show the Foundry playground + the hosted web app in a browser; narrate the admin-center publish steps from slide 8.",
        command="azd deploy  # then upload the Teams app package in Teams admin center",
    )


def s_rollout(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    heading(slide, "Rollout & operations", "Pilot, monitor, version, roll back", 10)
    cols = [
        ("PILOT", GREEN, ["Targeted group sideload/policy", "Validate SSO + approval on desktop, web, mobile", "UAT sign-off with evidence"]),
        ("SCALE", MSBLUE, ["Staged rollout via setup policy", "Monitor Teams usage reports", "Watch App Insights error rates"]),
        ("OPERATE", AMBER, ["Version bump in manifest per update", "Access reviews on the assigned group", "Owner + support contact recorded"]),
        ("ROLL BACK", RED, ["Set app status Blocked", "Remove setup-policy assignment", "Revert to prior package version"]),
    ]
    for index, (tag, accent, items) in enumerate(cols):
        x = 0.58 + index * 3.06
        rect(slide, x, 1.8, 2.85, 4.3, WHITE, line=LINE)
        rect(slide, x, 1.8, 2.85, 0.09, accent)
        text(slide, tag, x + 0.2, 2.05, 2.5, 0.3, size=15, font="Segoe UI Semibold", bold=True)
        for j, item in enumerate(items):
            rect(slide, x + 0.2, 2.62 + j * 0.95, 0.1, 0.1, accent)
            text(slide, item, x + 0.4, 2.52 + j * 0.95, 2.35, 0.85, size=11.5, value_color=MUTED)
    text(slide, "Decommission: remove from catalog, delete the Entra app, azd down the Foundry environment, revoke consents.",
         0.58, 6.3, 12.1, 0.3, size=12, value_color=INK, bold=True)
    add_notes(slide, "Operate it like any enterprise app: staged rollout, monitoring, versioning, a clear rollback, and a clean decommission path including azd down.")


def s_close(prs):
    closing_slide(
        prs,
        title_value="Governed AI,\ndelivered in Teams",
        rules=[
            ("01", "Reuse the hardened app as a tab"),
            ("02", "Two identities: user SSO, agent MI"),
            ("03", "Authorize every mutation endpoint"),
            ("04", "Publish and target through policy"),
        ],
        proof_headline="1 app",
        proof_caption="package, policy-governed",
        proof_items=["Teams tab + Entra SSO", "Foundry Hosted Agent backend", "Model Router contract", "Admin-center governance", "Auditable approval in Teams"],
        resources="learn.microsoft.com/microsoftteams · Teams admin center · aka.ms/teams-app-manifest · Foundry hosted agents",
    )


DECK = [
    s_title, s_goal, s_architecture, s_patterns, s_identity, s_hosting,
    s_package, s_publish, s_demo, s_rollout, s_close,
]


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Caldova in Microsoft Teams — Deploying the Hosted Agent"
    prs.core_properties.subject = "Teams tab, Entra SSO, Foundry Hosted Agent, admin-center governance"
    prs.core_properties.author = "Caldova Recall Control Tower"
    for builder in DECK:
        builder(prs)
    output = ROOT / "caldova-teams-deployment.pptx"
    prs.save(output)
    return output


if __name__ == "__main__":
    print(f"wrote {build().name}")
