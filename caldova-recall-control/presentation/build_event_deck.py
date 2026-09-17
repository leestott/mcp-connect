"""Build the 25-minute MCP Community Connect Bengaluru presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from deck_common import (
    EVENT_ART, SLIDE_H, SLIDE_W, add_notes, line, rect,
    text as shared_text,
)


INK = "20252A"
DEEP = "20252A"
WHITE = "FFFFFF"
MUTED = "58656D"
LINE = "D9E3E6"
MSBLUE = "006F78"
CYAN = "8BE2DC"
RED = "B73935"
AMBER = "F4C44E"
BLUE_SOFT = "EDF6F5"
DISPLAY = "Bahnschrift"


def text(slide, value, x, y, w, h, **kwargs):
    kwargs.setdefault("font", "Calibri")
    kwargs.setdefault("value_color", INK)
    return shared_text(slide, value, x, y, w, h, **kwargs)


def base(slide, dark=False, number=None):
    rect(slide, 0, 0, 13.333, 7.5, DEEP if dark else WHITE)
    rect(slide, 0, 0, 0.13, 7.5, CYAN if dark else MSBLUE)
    text(slide, "MCP COMMUNITY CONNECT", 0.55, 7.12, 4.3, 0.2,
         size=9, value_color=CYAN if dark else MSBLUE, bold=True)
    text(slide, "BENGALURU  /  26 SEP 2026", 8.2, 7.12, 3.9, 0.2,
         size=9, value_color=CYAN if dark else MUTED, align=PP_ALIGN.RIGHT)
    if number is not None:
        text(slide, f"{number:02d}", 12.35, 7.06, 0.45, 0.3,
             size=13, font=DISPLAY, value_color=CYAN if dark else MSBLUE)


ROOT = Path(__file__).resolve().parent
EVENT_URL = "https://globalai.community/e/bd1o37ln"
VIDEO_URL = "https://www.youtube.com/watch?v=uydwDk91Y9Y"
SESSION_TITLE = (
    "From Prototype to Production: Engineering Agent Systems with MCP, "
    "Multi-Agent Patterns and Real Work"
)


def new_slide(prs, kicker, title_value, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, dark=dark, number=len(prs.slides))
    text(slide, kicker.upper(), 0.65, 0.42, 12.0, 0.28,
        size=11, value_color=CYAN if dark else MSBLUE, bold=True)
    text(slide, title_value, 0.62, 0.86, 12.1, 0.85,
        size=32, font=DISPLAY, bold=True, value_color=WHITE if dark else INK)
    return slide


def block(slide, title_value, detail, x, y, w, *, accent=MSBLUE, dark=False):
    rect(slide, x, y, w, 1.02, "30393E" if dark else BLUE_SOFT)
    rect(slide, x, y, 0.055, 1.02, accent)
    text(slide, title_value, x + 0.16, y + 0.13, w - 0.3, 0.34,
        size=18, font=DISPLAY, bold=True, value_color=WHITE if dark else INK)
    text(slide, detail, x + 0.16, y + 0.58, w - 0.3, 0.3,
        size=12, value_color=CYAN if dark else MUTED)


def arrow(slide, x, y, width=0.4, accent=MSBLUE):
    line(slide, x, y, x + width, y, accent, 1.8)
    line(slide, x + width - 0.1, y - 0.08, x + width, y, accent, 1.8)
    line(slide, x + width - 0.1, y + 0.08, x + width, y, accent, 1.8)


def picture_fit(slide, filename, x, y, w, h):
    picture = slide.shapes.add_picture(str(ROOT / "assets" / filename), Inches(x), Inches(y))
    ratio = min(Inches(w) / picture.width, Inches(h) / picture.height)
    picture.width = int(picture.width * ratio)
    picture.height = int(picture.height * ratio)
    return picture


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, number=1)
    text(slide, "MCP Community Connect", 0.65, 0.55, 9.4, 0.6,
        size=28, font=DISPLAY, bold=True, value_color=MSBLUE)
    slide.shapes.add_picture(str(EVENT_ART), Inches(10.7), Inches(0.34), width=Inches(1.8))
    text(slide, "From Prototype", 0.6, 1.95, 12.0, 1.0,
        size=58, font=DISPLAY, bold=True)
    text(slide, "to Production", 0.6, 3.0, 12.0, 1.0,
        size=58, font=DISPLAY, bold=True, value_color=MSBLUE)
    text(slide, "Engineering Agent Systems with MCP,\nMulti-Agent Patterns and Real Work",
        0.65, 4.35, 11.3, 0.95, size=25)
    line(slide, 0.65, 5.64, 12.6, 5.64, LINE)
    text(slide, "Lee Stott", 0.65, 5.98, 4.0, 0.42, size=22, font=DISPLAY, bold=True)
    text(slide, "Principal Cloud Advocate Manager, Microsoft", 0.65, 6.45, 8.4, 0.28,
        size=14, value_color=MUTED)
    text(slide, "11:00-11:25 IST", 9.1, 6.08, 3.5, 0.4,
        size=20, value_color=MSBLUE, align=PP_ALIGN.RIGHT)
    add_notes(slide, f"""
TIME: 0:00-0:45. Open, establish stakes, and move on.

SAY: "Connecting an agent to a tool is straightforward. Knowing whether that tool
should be called, what happens if it fails, and who is accountable for the result
is the engineering problem. Today our system will assess a recall affecting 2,196
units. The question is not whether a model can write a convincing recommendation.
It is whether the application can act safely, explain what happened and survive a retry."

POINT: Introduce the session as {SESSION_TITLE}. This is a developer session about
boundaries, orchestration and evidence, not prompt tricks or a healthcare product pitch.
Microsoft Agent Framework supplies the workflow implementation; Microsoft Foundry is
the hosted example; MCP supplies the tool boundary. These have different responsibilities.

OPTIONAL DEPTH: The audience does not need protocol expertise. Explain MCP as a shared
language through which an agent application discovers and calls external capabilities.
The rest of the talk is about the software around that connection. The patterns apply
with other hosts and model providers as well.

CAUTION: Caldova and its operational data are fictional. We will show a local controls
demo and a separate hosted reasoning demo, not a deployed enterprise recall product.
Some patterns are compared on slides rather than executed. Do not promise dynamic
supervisor routing in the live example.

TRANSITION: "First, meet the workload. What exactly would a correct outcome look like?"
Source: {EVENT_URL}
""")


def scenario(prs):
    slide = new_slide(prs, "Meet Caldova | synthetic pharmaceutical recall", "A decision brief is not permission to act.")
    text(slide, "2,196", 0.55, 2.0, 4.0, 1.3, size=79, font=DISPLAY, bold=True, value_color=RED)
    text(slide, "units need a decision", 0.65, 3.45, 4.1, 0.55, size=24)
    text(slide, "4 locations\n1 named approval\n0 duplicate mutations", 0.65, 4.35,
        4.0, 1.6, size=25, font=DISPLAY)
    picture_fit(slide, "control-tower-initial.png", 4.8, 2.0, 7.85, 4.65)
    text(slide, "B-2408-AX7  |  Saved local Control Tower view", 4.8, 6.58, 7.7, 0.25,
        size=12, value_color=MUTED)
    add_notes(slide, """
TIME: 0:45-1:30. Establish the acceptance contract before the architecture.

SAY: "This is Caldova Recall Control Tower. A fictional cold-chain incident affects
one batch across four stock locations. The system has to establish recall facts,
calculate exposure, inspect supplier constraints and prepare a decision brief. A
named person must approve before the application changes stock. Retrying must not
change the same positions again. That is our definition of success."

POINT: Start at 2,196 units, then point to the screenshot of the actual local UI.
The large number is an operational invariant, not a model-generated estimate.
The fixture contains Bengaluru DC 1,240, Chennai DC 760, Bengaluru Store 031 112,
and Mysuru Store 014 84. The batch is B-2408-AX7, Caldova Relief 20 mg tablets.
Do not read all locations aloud unless asked; they are useful when checking a live answer.

OPTIONAL DEPTH: This is beyond a chatbot because a state transition matters. A fluent
answer can still omit evidence, invent stock or claim an action that never occurred.
We therefore distinguish the decision brief from permission, the mutation from the
description of a mutation, and an audit record from conversational text.

CAUTION: This is a saved screenshot. These counts are synthetic acceptance criteria,
not evidence that a fresh invocation or tests passed today. Avoid saying the demo
is medically validated or connected to real inventory.

TRANSITION: "Now let us place the UI, agents and protocol in the solution architecture."
""")


def protocol(prs):
    slide = new_slide(prs, "MCP tool design | actual server contract, selected fields", "Discover the schema. Inspect the result.")
    text(slide, "initialize  >  tools/list  >  tools/call  >  result", 0.65, 1.94,
            12.0, 0.42, size=23, font=DISPLAY, value_color=MSBLUE, bold=True)
    line(slide, 6.35, 2.65, 6.35, 6.1, LINE)
    text(slide, "DISCOVERY  /  locate_inventory", 0.65, 2.65, 5.45, 0.36,
            size=18, font=DISPLAY, bold=True)
    text(slide, 'batch_id: required string\nminLength: 1\nmaxLength: 64\npattern: ^[A-Z0-9-]+$',
            0.65, 3.25, 5.45, 1.6, size=20, font="Cascadia Code")
    text(slide, "readOnlyHint: true", 0.65, 5.09, 5.45, 0.35,
            size=18, font="Cascadia Code", value_color=MSBLUE)
    text(slide, "outputSchema is a generic object.\nStructured does not mean fact-checked.",
            0.65, 5.65, 5.45, 0.8, size=18, value_color=MUTED)
    text(slide, "CALL  /  params excerpt", 6.85, 2.65, 5.75, 0.36,
            size=18, font=DISPLAY, bold=True)
    text(slide, '{"name": "locate_inventory",\n "arguments": {\n   "batch_id": "B-2408-AX7"\n }}',
            6.85, 3.25, 5.75, 1.5, size=18, font="Cascadia Code")
    text(slide, "RESULT  /  structuredContent excerpt", 6.85, 4.99, 5.75, 0.36,
            size=17, font=DISPLAY, bold=True)
    text(slide, '{"total_units": 2196,\n "locations": 4}',
            6.85, 5.55, 5.75, 0.8, size=20, font="Cascadia Code", value_color=MSBLUE)
    text(slide, "A model may propose a call. The client sends it. Server policy decides whether it may act.",
            0.65, 6.62, 12.0, 0.34, size=16, value_color=RED, bold=True)
    add_notes(slide, """
TIME: 3:00-4:00. Teach tool design using one concrete contract.

SAY: "A useful tool is narrower than the business request. Instead of a generic
execute-anything endpoint, locate_inventory answers one question: where is this
batch and how much stock is affected? Its input identifies the batch. Its result
contains positions, units and a location count that our workflow can verify."

POINT: Read the lifecycle across the top. Initialization negotiates protocol version
and capabilities; the SDK manages that exchange and the initialized notification.
tools/list describes available tools, including inputSchema and optional outputSchema.
tools/call names one tool and supplies arguments. The response can contain content
blocks, structuredContent, and isError. These examples are selected fields, not full
JSON-RPC envelopes. Do not imply that a model opens the connection itself.

TECHNICAL WALKTHROUGH: The required batch_id is a string of 1 to 64 characters,
restricted to uppercase letters, digits and hyphens. A JSON request can be syntactically
valid yet fail this schema. A request can also pass its schema and fail business policy.
The live inspection demonstrates both. In the Python SDK the wire fields inputSchema,
structuredContent and isError appear as input_schema, structured_content and is_error.

IMPORTANT LIMIT: The real return annotation is dict[str, Any], so this server advertises
a generic object outputSchema with additionalProperties allowed. It does not enforce
required inventory fields or validate the sum. The client/test asserts 2,196 and four
locations. A typed result model plus explicit invariant checks would strengthen the
contract; that improvement is not implemented here. Structured data is easier to check,
but merely receiving structured JSON is not evidence of factual correctness.

OPTIONAL DEPTH: Tool descriptions should say when to use the tool, its input constraints,
side effects and failure behavior. Prefer structured results that downstream code can
check. Bound result size and avoid returning an entire operational database to a model.
Read-only or destructive annotations help a client reason about intent, but neither
an annotation nor a prompt grants authorization. Treat returned prose as untrusted data.

CAUTION: The agent-facing bridge in this implementation allow-lists read tools.
Remote authentication and action authorization remain separate requirements. MCP
does not choose the workflow graph or implement the inventory transaction for us.

TRANSITION: "The tool contract is clear. Which orchestration shape should consume it?"
Reference: https://modelcontextprotocol.io/docs/learn/architecture
""")


def architecture(prs):
    slide = new_slide(prs, "Solution architecture | two separate execution paths", "One domain. Two demos. Explicit boundaries.")
    text(slide, "01  LOCAL CONTROL PATH", 0.65, 1.93, 6.5, 0.3, size=13, value_color=MSBLUE, bold=True)
    for horizontal, title_value, detail in [
        (0.65, "Control Tower", "Browser + FastAPI"),
        (3.77, "MCP client", "In-memory tool calls"),
        (6.89, "MCP server", "Typed tool contract"),
        (10.01, "Policy + store", "Approval / replay / audit"),
    ]:
        block(slide, title_value, detail, horizontal, 2.44, 2.68)
    for horizontal in [3.34, 6.46, 9.58]:
        arrow(slide, horizontal, 2.95, 0.34)
    text(slide, "Human approves in the application. No LLM call in this UI path.",
         0.65, 3.68, 12.0, 0.37, size=17, value_color=RED, bold=True)
    text(slide, "02  HOSTED REASONING PATH  |  MICROSOFT FOUNDRY", 0.65, 4.26,
         12.0, 0.3, size=13, value_color=MSBLUE, bold=True)
    for horizontal, title_value, detail in [
        (0.65, "Responses host", "Separate hosted invocation"),
        (3.77, "Agent Framework", "4 agents + model client"),
        (6.89, "CLI / JSON bridge", "Read-only allowlist"),
        (10.01, "In-process MCP", "Isolated interpreter + fixtures"),
    ]:
        block(slide, title_value, detail, horizontal, 4.77, 2.68)
    for horizontal in [3.34, 6.46, 9.58]:
        arrow(slide, horizontal, 5.28, 0.34)
    text(slide, "Shared domain code and synthetic data, not a shared live store or a wired UI-to-agent call.",
         0.65, 6.3, 12.0, 0.53, size=17, value_color=MUTED)
    add_notes(slide, """
TIME: 1:30-3:00. This diagram is the anchor for the rest of the session.

SAY: "There are two separate paths here. At the top, the browser calls the local
FastAPI application. That application makes real in-memory MCP calls to the server.
The domain layer owns approval validation, inventory changes, replay behavior and
audit. This path demonstrates controls deterministically; it does not invoke an LLM."

POINT: Trace the top lane left to right. The human approval arrives through the
application. The application retains the usable approval token and exposes a local
handle to the UI. A typed approver name is demonstration input, not authenticated identity.

SAY: "Below is the actual agent path. A separate request reaches a Responses host
in Microsoft Foundry. Microsoft Agent Framework runs four agents with a model client.
Their tools call a read-only bridge, which uses an isolated MCP interpreter. The
agents can inspect recall evidence, but they cannot approve or quarantine stock."

OPTIONAL DEPTH: These paths reuse domain code and synthetic fixtures, not a shared
live database. The current Control Tower does not call the hosted workflow. A
deployment of the same code is not evidence that local mutation state appears in
the hosted agent. Wiring a production UI, identity boundary and durable store would
be additional work. The diagram describes logical components, not network segmentation.

HOST, CLIENT, SERVER: An MCP host is the application that owns user interaction and
tool-use policy. A client is its protocol connection/session; a server exposes capabilities.
The model is not the MCP client. In this local path FastAPI owns Client(mcp). The hosted
agent gets explicitly registered Python wrappers, not dynamically discovered MCP tools.
main.py starts mcp_v2_bridge.py with a tool name and JSON argument on the command line;
that child opens Client(mcp) in process and prints JSON back. This outer CLI/JSON adapter
is custom application code, not an MCP transport. The new inspect_mcp.py demonstration
instead launches caldova_mcp.py using StdioServerParameters and genuinely exercises
MCP stdio. Distinguish the three paths when pointing at the diagram.

CAUTION: Do not draw an imaginary UI-to-Foundry integration while narrating. Hosting,
model access, MCP calls and business authorization are separate responsibilities.
Source anchors: control_tower_api.py run_analysis/invoke; main.py build_workflow;
mcp_v2_bridge.py read-tool boundary.

TRANSITION: "MCP is the contract inside both paths. Here is what a well-designed tool exposes."
""")


def patterns(prs):
    slide = new_slide(prs, "Choose the orchestration pattern", "Use the smallest workflow that can do the job.")
    for index, (title_value, purpose, status) in enumerate([
        ("Sequential", "Known dependencies", "IMPLEMENTED"),
        ("Supervisor routing", "Choose the next specialist", "COMPARISON"),
        ("Parallel collaboration", "Independent evidence, then join", "COMPARISON"),
    ]):
        horizontal = 0.65 + index * 4.15
        text(slide, status, horizontal, 2.0, 3.8, 0.26, size=11,
             value_color=MSBLUE if index == 0 else MUTED, bold=True)
        text(slide, title_value, horizontal, 2.45, 3.8, 0.85, size=25, font=DISPLAY, bold=True)
        text(slide, purpose, horizontal, 5.4, 3.7, 0.65, size=18, value_color=MUTED)
        if index < 2:
            line(slide, horizontal + 3.88, 2.0, horizontal + 3.88, 6.0, LINE)
    for index, label_value in enumerate(["Read", "Assess", "Brief"]):
        horizontal = 0.65 + index * 1.22
        rect(slide, horizontal, 3.85, 1.02, 0.72, BLUE_SOFT)
        text(slide, label_value, horizontal + 0.07, 4.02, 0.88, 0.28, size=13, bold=True)
        if index < 2:
            arrow(slide, horizontal + 1.03, 4.21, 0.18)
    block(slide, "Router", "Bounded choices", 4.8, 3.4, 1.8)
    for vertical, label_value in [(3.43, "Stock"), (4.47, "Supplier")]:
        line(slide, 6.6, 3.91, 6.92, vertical + 0.3, MSBLUE, 1.5)
        rect(slide, 6.92, vertical, 1.45, 0.62, BLUE_SOFT)
        text(slide, label_value, 7.05, vertical + 0.15, 1.17, 0.3, size=14)
    for vertical, label_value in [(3.43, "Stock"), (4.47, "Supplier")]:
        rect(slide, 9.15, vertical, 1.47, 0.62, BLUE_SOFT)
        text(slide, label_value, 9.25, vertical + 0.15, 1.27, 0.3, size=14)
        line(slide, 10.62, vertical + 0.3, 11.17, 4.21, MSBLUE, 1.5)
    block(slide, "Join", "Reconcile", 11.17, 3.7, 1.47)
    text(slide, "Human approval is an authority boundary across every pattern, not another agent vote.",
         0.65, 6.4, 12.0, 0.45, size=18, value_color=RED, bold=True)
    add_notes(slide, """
TIME: 4:00-5:15. Compare patterns without promising three live implementations.

SAY: "Start with the smallest workflow that can do the job. If every step depends
on evidence from the previous one, a sequential workflow gives you a predictable
order and a clear place to validate the handoff. That is what this sample implements."

POINT: Move to the middle diagram. "Supervisor routing is different. A controller
chooses which specialist should handle the next task. That can help when requests
vary, but it introduces routing errors, loops and a termination problem. Bound the
choices, give the router explicit completion criteria and evaluate its decisions."

SAY: "Use parallel collaboration when evidence sources are independent. Fan out,
then join and reconcile the results. You gain potential latency benefits but need
a partial-failure policy, consistent inputs and rules for disagreement. More agents
are not automatically more intelligence. Sometimes one tool-enabled agent is sufficient."

OPTIONAL DEPTH: Supervisor routing and decentralized handoffs are not identical:
one centralizes the choice, the other delegates it between participants. Neither is
shown live here. In our sample the agent named supervisor is the final synthesizer,
not a dynamic router. Human approval can sit after any of these patterns because
business authority is independent of who produced the recommendation.

CAUTION: Only sequential collaboration and application-level human approval are
demonstrated. Routing and parallel execution are design comparisons. If the abstract
must promise all patterns live, that requires new demo implementation and rehearsal.

TRANSITION: "Let us look at the exact sequence our Agent Framework code builds."
Reference: https://github.com/microsoft/agent-framework/tree/main/python/packages/orchestrations
""")


def agents(prs):
    slide = new_slide(prs, "Microsoft Agent Framework | implemented sequence", "Four specialists. Three tools. No mutation rights.", dark=True)
    for index, (name, tool_name, output) in enumerate([
        ("Recall Triage", "get_recall_notice", "Validated recall"),
        ("Inventory", "locate_inventory", "Exposure + positions"),
        ("Supplier", "get_supplier_status", "Constraints + unknowns"),
        ("Supervisor", "NO TOOLS", "Decision brief"),
    ]):
        horizontal = 0.65 + index * 3.14
        text(slide, f"0{index + 1}", horizontal, 2.23, 2.65, 0.75,
             size=42, font=DISPLAY, value_color=CYAN)
        text(slide, name, horizontal, 3.29, 2.7, 0.5, size=25, font=DISPLAY, value_color=WHITE, bold=True)
        text(slide, tool_name, horizontal, 4.03, 2.7, 0.5, size=13, font="Cascadia Code", value_color=CYAN)
        text(slide, output, horizontal, 4.8, 2.7, 0.65, size=18, value_color=WHITE)
        if index < 3:
            arrow(slide, horizontal + 2.73, 3.52, 0.28, accent=CYAN)
    line(slide, 0.65, 5.91, 12.65, 5.91, "58656D")
    text(slide, "Fixed WorkflowBuilder edges. The supervisor synthesizes; it does not route.",
         0.65, 6.2, 12.0, 0.5, size=19, value_color=CYAN)
    add_notes(slide, """
TIME: 5:15-6:00. Explain the implemented graph; keep code details for questions.

SAY: "Recall Triage establishes the batch. Inventory Impact measures exposure.
Supplier and Compliance identifies constraints and unknowns. The supervisor receives
the evidence and produces a brief. Each specialist has one read tool. The supervisor
has no tools. Its job is synthesis, not acquiring new authority."

POINT: Trace the arrows. The actual WorkflowBuilder has three fixed edges:
triage to inventory, inventory to compliance, compliance to supervisor. The workflow
returns output from the supervisor. This is collaboration through shared evidence,
not a free-form conversation in which agents negotiate what to do next.

OPTIONAL DEPTH: Triage uses last-agent context; downstream executors use full context.
The final agent also sees the original question so it can distinguish a focused fact
request from a broad assessment. More context helps preserve facts, but increases
token use and does not guarantee complete synthesis. Evaluate the final brief against
the actual tool results, not against whether it sounds comprehensive.

CAUTION: Instructions to call a tool exactly once guide model behavior; they are not
a transaction or authorization mechanism. The bridge's read-only restriction is code.
Neither request_approval nor quarantine_batch is exposed to these agents. Do not call
the final synthesis step dynamic supervisor routing.

TRANSITION: "Before trusting a model with the evidence, prove that the controls work.
We will start with the local MCP and policy path."
""")


def local_demo(prs):
    slide = new_slide(prs, "Demo 01 | local MCP + deterministic controls | 5 minutes", "Can it fail safely, then recover?", dark=True)
    picture_fit(slide, "control-tower-initial.png", 0.65, 2.05, 8.0, 4.6)
    for index, (title_value, detail) in enumerate([
        ("DISCOVER + CALL", "Real stdio MCP inspection"),
        ("REJECT + VERIFY", "Bad input; denied write"),
        ("APPROVE + ACT", "Local browser + MCP"),
        ("REPLAY", "Zero further changes"),
    ]):
        vertical = 2.1 + index * 1.12
        text(slide, title_value, 9.0, vertical, 3.6, 0.38, size=21, font=DISPLAY, value_color=CYAN, bold=True)
        text(slide, detail, 9.0, vertical + 0.49, 3.6, 0.42, size=17, value_color=WHITE)
    text(slide, "Local UI: deterministic stage display, not a model invocation. Image: saved starting state.",
         0.65, 6.7, 12.0, 0.24, size=11, value_color=CYAN)
    add_notes(slide, """
TIME: 6:00-11:00. Five minutes; stop on time even if a step is incomplete.

SAY: "This first run is deliberately deterministic. The UI displays four stages,
but the local endpoint performs three real MCP reads and assembles the brief in code.
We are testing tool access and state transitions here, not claiming an LLM executed."

DEMO CUES:
0. At 6:00 run the prepared terminal command from the repository root:
./caldova-recall-control/.venv/Scripts/python.exe caldova-recall-control/scripts/inspect_mcp.py
Budget 75 seconds. It uses the installed MCP SDK to launch the real stdio server.
Point to CONNECT and the negotiated version, then DISCOVER and seven server tools.
Read the inventory inputSchema constraints and generic outputSchema. Show the actual
tools/call arguments and structured result, including four positions and 2,196 units.
The malformed batch produces isError true. A schema-valid but unapproved quarantine
also produces isError true. A follow-up read verifies unchanged inventory. No model
or cloud credentials are involved, and this child does not change browser state.
Do not mistake discovery of seven tools for permission to expose all seven to an agent.
The bridge exposes only three read tools; request_approval and reset_demo are powerful
application/demo capabilities and would need a real identity boundary in production.
1. Switch to the already running http://127.0.0.1:8091 and reset. Show the active
recall and expected 2,196 units across four positions. Do not rely on a particular
status badge: the UI may say ASSESSMENT before analysis.
2. Select Run analysis. Point to the three tool names, input batch, outcomes and
durations. The completed supervisor stage is a deterministic summary in this path.
3. For denial, use a prepared PowerShell 7 terminal. There is no unauthenticated
quarantine button in the UI. This request uses a deliberately invalid local handle:
Invoke-WebRequest http://127.0.0.1:8091/api/quarantine -Method POST -ContentType 'application/json' -Body '{"approval_id":"not-approved"}' -SkipHttpErrorCheck
Expect HTTP 403 and "Valid named approval is required". Never substitute a real token.
4. Select Approve quarantine. Enter Asha Rao, Responsible Pharmacist. Submitting the
dialog records the approval and immediately invokes quarantine; there is no separate
second quarantine button. Check four positions changed and 2,196 units processed.
5. Select Replay quarantine. Verify positions_changed is zero and inspect the audit.

EXPLAIN: Ask the audience to distinguish two denials. The terminal receives an MCP
CallToolResult with isError; the browser API returns HTTP 403 before its invalid handle
can authorize the call. HTTP 403 is not an MCP protocol error. The terminal deliberately
uses a schema-valid invalid credential so policy, not the regex, rejects the operation.
Skip the extra HTTP denial on stage if the terminal segment runs long.
A disabled button alone is not authorization. The direct API denial proves
the application checks the handle. However, entering a name is still not enterprise
authentication. Keep the usable credential out of the browser display and audit.

RECOVERY: If the server or action fails, do not improvise setup on stage. Move to
slide 8, label it a saved screenshot, and explain the expected transitions. Do not
claim a fresh run passed. The replay guarantee is local, not yet cross-replica proof.

TRANSITION: "What matters is the state change, not the confidence of the explanation."
""")


def evidence(prs):
    slide = new_slide(prs, "Demo fallback | saved screenshot", "The proof is in the state change")
    picture_fit(slide, "control-tower-quarantined.png", 0.65, 2.0, 8.0, 4.65)
    for index, (value, caption) in enumerate([
        ("2,196", "units on the first call"),
        ("0", "additional changes on replay"),
        ("Audit", "named approval; no credential"),
    ]):
        position = 1.95 + index * 1.52
        text(slide, value, 9.0, position, 3.65, 0.7,
             size=38, value_color=MSBLUE, bold=True)
        text(slide, caption, 9.0, position + 0.77, 3.65, 0.55,
             size=16, value_color=MUTED)
    add_notes(slide, """
TIME: 11:00-11:30. A brief debrief, or a clearly labelled fallback.

SAY: "There are three things to verify. The authorized first call processes the
affected stock. A replay changes no additional positions. The audit records the
decision without exposing a credential somebody could reuse. Those are properties
of the application and domain controls, not of how persuasive the brief sounds."

POINT: This image is a saved quarantined-state screenshot. The replay number at the
right summarizes the acceptance criterion; it is not proof that this image alone
captures every step. When the live demo succeeds, reference the actual first and
replay responses. When it fails, say plainly that the screenshot is the fallback.

OPTIONAL DEPTH: Idempotency means retrying an operation does not create an additional
business effect. It does not necessarily mean every response field is identical.
A replay can report the same batch outcome while returning zero newly changed positions.
In distributed systems, the idempotency decision must be stored atomically with the
business mutation. A process-local dictionary is not sufficient across replicas.

CAUTION: A named approval is a demo record, not evidence of verified professional
identity. The screenshot does not prove today's tests, a hosted invocation or a
production deployment passed. Keep these claims separate during questions.

TRANSITION: "A correct state change is one half of reliability. When the answer is
wrong or slow, we also need to find the exact point where evidence diverged."
""")


def debugging(prs):
    slide = new_slide(prs, "Observability + debugging | schematic, not captured telemetry", "Find the first point where evidence goes wrong.")
    labels = ["Request", "Triage", "Inventory", "Supplier", "Brief"]
    for index, label_value in enumerate(labels):
        vertical = 2.12 + index * 0.66
        text(slide, label_value, 0.65, vertical + 0.05, 1.55, 0.35, size=17)
        rect(slide, 2.25 + index * 0.59, vertical, 4.9 - index * 0.58, 0.4,
             RED if index == 4 else MSBLUE)
    text(slide, "Relative sequence only; bar lengths are not measurements.", 0.65, 5.76,
         6.7, 0.58, size=13, value_color=MUTED)
    for index, (label_value, detail) in enumerate([
        ("WRONG TOOL RESULT?", "Inspect arguments, schema and source."),
        ("RIGHT RESULT, WRONG BRIEF?", "Inspect context handoff and synthesis."),
        ("SLOW OR FAILED CALL?", "Inspect duration, timeout and cleanup."),
    ]):
        vertical = 2.08 + index * 1.42
        text(slide, label_value, 8.0, vertical, 4.6, 0.34, size=13, value_color=MSBLUE, bold=True)
        text(slide, detail, 8.0, vertical + 0.48, 4.6, 0.7, size=21)
    text(slide, "Correlate decisions and policy outcomes. Never log usable approval credentials.",
         0.65, 6.5, 12.0, 0.4, size=18, value_color=RED, bold=True)
    add_notes(slide, """
TIME: 11:30-13:00. Teach a debugging method before opening a hosted trace.

SAY: "Start with the first incorrect piece of evidence, not the last sentence in
the response. If the inventory tool returned the wrong result, inspect its arguments,
validation and data source. If the tool result is correct but the final brief is
wrong, inspect the context passed between steps and the supervisor's synthesis."

POINT: Walk down request, triage, inventory, supplier and brief. This is an illustrative
sequence, not a captured trace or performance chart. The final red bar marks where
an answer could become incorrect, not an observed incident. Bar lengths are not timings.
For a slow or failed call, inspect duration, timeouts, process exit and cleanup before
changing prompts. A timeout is an infrastructure symptom, not a missing instruction.

OPTIONAL DEPTH: The local API records tool, filtered input, correlation ID, outcome
and duration. The hosted bridge bounds concurrent child processes, times out communication,
reaps failed children and returns controlled errors. Those are inspectable code facts.
End-to-end hosted trace propagation must be verified in a real run; do not assume a
correlation ID automatically crosses every boundary.

DEMO CUE: In the next demo find one tool result and compare its facts with the final
brief. If a supplier fact is absent, determine whether it was missing at source or
lost during synthesis. A useful evaluation checks coverage as well as refusal behavior.

CAUTION: Do not expose approval tokens, raw secrets or sensitive operational payloads
while tracing. Observability needs enough evidence to debug, not indiscriminate logging.

TRANSITION: "Now we will apply that method to the real Agent Framework workflow."
""")


def production(prs):
    slide = new_slide(prs, "Identity / security / governance", "Three boundaries hosting does not remove.")
    for index, (title_value, subtitle, detail, accent) in enumerate([
        ("PERSON", "Who may approve?", "Authenticate the operator.\nAuthorize the exact action.\nRecord the decision.", MSBLUE),
        ("WORKLOAD", "What may the agent access?", "Use a runtime identity.\nGrant narrow read access.\nKeep deployment rights separate.", INK),
        ("STATE", "What survives a retry?", "Persist approvals and audit.\nMake mutations atomic.\nPreserve replay protection.", RED),
    ]):
        horizontal = 0.65 + index * 4.15
        rect(slide, horizontal, 2.1, 3.78, 0.07, accent)
        text(slide, title_value, horizontal, 2.55, 3.78, 0.62, size=33, font=DISPLAY, value_color=accent, bold=True)
        text(slide, subtitle, horizontal, 3.57, 3.78, 0.9, size=23, font=DISPLAY, bold=True)
        text(slide, detail, horizontal, 4.85, 3.78, 1.45, size=20)
    text(slide, "Release requirements, not a claim that the localhost demo is enterprise-ready.",
         0.65, 6.65, 12.0, 0.28, size=13, value_color=MUTED)
    add_notes(slide, """
TIME: 17:00-18:30. Separate human identity, workload access and durable state.

SAY: "Moving the process into a managed host does not finish the security design.
The person, the workload and the state each need a boundary. A person must be
authenticated and authorized to approve this action. The agent's workload identity
needs only the resources its job requires. And the records that make approval and
replay meaningful must survive beyond the lifetime of one process."

POINT: In PERSON, distinguish a typed name from verified identity. In WORKLOAD,
separate deployment permissions from runtime access. In STATE, ask what happens
when the response is lost, the process restarts or a second replica receives the retry.
The API must enforce authorization even when the normal UI flow is bypassed.

OPTIONAL DEPTH: Governance includes who can publish tools, who reviews changed schemas
or permissions, which agent version is approved, what evidence supports promotion,
how long audit is retained and how a release is rolled back. Record the decision
and policy outcome without storing a usable credential in a trace or transcript.
For destructive operations, establish transaction and idempotency rules at the store.

CAUTION: These are production requirements, not completed sample features. The local
API has no enterprise caller authentication, and its state is in memory. Do not expose
it as a public production service or claim horizontal scaling is safe because the
host can start more containers. Infrastructure scale and correctness are different.

TRANSITION: "The wider MCP ecosystem can improve the experience, but none of those
new capabilities removes these three responsibilities."
""")


def hosted_demo(prs):
    slide = new_slide(prs, "Demo 02 | Microsoft Foundry + Agent Framework | 4 minutes", "Now run the actual agent workflow.", dark=True)
    text(slide, "ASSESS", 0.65, 2.18, 7.0, 0.85, size=52, font=DISPLAY, value_color=CYAN, bold=True)
    text(slide, "Batch B-2408-AX7.\nReport exposure, supplier status\nand required human action.",
         0.65, 3.43, 7.25, 1.7, size=29, value_color=WHITE)
    for index, (title_value, detail) in enumerate([
        ("VERSION", "Identify the running deployment"),
        ("WORKFLOW", "Inspect sequence and tool outputs"),
        ("TRACE", "Follow evidence into the final brief"),
        ("BOUNDARY", "No approval or mutation by agents"),
    ]):
        vertical = 2.13 + index * 1.09
        text(slide, title_value, 8.95, vertical, 3.65, 0.32, size=15, value_color=CYAN, bold=True)
        text(slide, detail, 8.95, vertical + 0.42, 3.65, 0.63, size=19, value_color=WHITE)
    text(slide, "Fallback: show source + a saved trace labelled with its version and date; never imply a live run.",
         0.65, 6.63, 12.0, 0.34, size=12, value_color=CYAN)
    add_notes(slide, """
TIME: 13:00-17:00. Four minutes; all tabs and authentication must be ready beforehand.

SAY: "This is the actual agent execution. We are invoking the deployed Responses
host separately from the Control Tower. Microsoft Agent Framework coordinates the
specialists, the model produces their reasoning, and their allowed tools cross the
read-only MCP bridge. The same fixture makes the facts comparable with the local demo."

DEMO CUES:
1. Show the running agent version and runtime identity from the deployment view.
Do not quote an old version number from memory or conflate the deployment principal
with the runtime identity. Show only identifiers appropriate for the audience.
2. Submit the prepared broad recall assessment for B-2408-AX7. Check the response
against four locations and 2,196 units. Look for supplier status, uncertainty and
required human action, not just a fluent first paragraph.
3. Open the matching trace if available. Locate workflow steps, at least one tool
input/result and the final brief. Explain one observed duration without turning it
into an unsupported benchmark. Verify which spans actually exist in this deployment.
4. Point out that none of these agents has approval or quarantine tools. A request
for action should describe the required application approval, not claim stock changed.

OPTIONAL DEPTH: The final supervisor synthesizes fixed upstream steps. It is not
making dynamic routing decisions. The hosted fixture does not inherit a local UI
quarantine: the process state and invocation path are separate.

RECOVERY: If invocation stalls, use only a prepared saved trace labelled with date
and version. If no trace is available, show build_workflow and the read-tool wrappers
as source evidence and explicitly say live execution was unavailable. Never invent
trace spans, identity verification, successful routing or evaluation results.

TRANSITION: "Hosting made execution manageable. Here are the boundaries we must
still engineer before calling the application production-ready."
""")


def mcp_live(prs):
    slide = new_slide(prs, "MCP Live! | next design questions", "Extend the experience. Keep the controls.")
    for index, (title_value, chapter, question, accent) in enumerate([
        ("Toolboxes", "01:20:23", "Which tools should a workload discover?", MSBLUE),
        ("MCP auth", "02:32:29", "Which caller may act on this resource?", INK),
        ("MCP Apps", "03:21:05", "Can an embedded UI improve evidence review?", RED),
        ("Event-driven agents", "03:45:08", "Can recall arrival trigger analysis, not approval?", MSBLUE),
    ]):
        horizontal = 0.65 + (index % 2) * 6.25
        vertical = 2.05 + (index // 2) * 2.18
        text(slide, chapter, horizontal, vertical, 5.6, 0.25, size=12, value_color=accent, bold=True)
        text(slide, title_value, horizontal, vertical + 0.46, 5.6, 0.52, size=27, font=DISPLAY, bold=True)
        text(slide, question, horizontal, vertical + 1.15, 5.6, 0.8, size=21, value_color=MUTED)
    text(slide, "Extension ideas, not implemented features of this demo.",
         1.35, 6.65, 11.2, 0.28, size=13, value_color=RED, bold=True)
    add_notes(slide, f"""
TIME: 18:30-19:30. One-minute outlook, not a second product tour.

SAY: "MCP Live highlights where the ecosystem is expanding. We can turn those topics
into design questions for this workload. As tool access grows, how do we curate what
each workload discovers? For a remote server, which caller is allowed to act? Could
an embedded interface make evidence review clearer? Could a recall event start the
analysis without also authorizing a mutation?"

POINT: Read the questions, not the timestamps. Toolboxes connects to governed tool
access. Authentication connects to caller identity. MCP Apps connects to interactive
evidence review. Event-driven agents connects to initiating work when the business
changes rather than waiting for a chat message. The design principle stays constant:
an improved interface or trigger does not replace server-side authority.

OPTIONAL DEPTH: These are extension ideas, not features of Caldova. An event-triggered
analysis could be useful, but it would still need deduplication, bounded retries,
failure visibility and a human approval boundary. A Teams tab is not an MCP App.
Do not turn this slide into unverified setup guidance or a claim of interoperability.

ATTRIBUTION: Topics are from the published video description, not a full transcript
review. The questions are our engineering interpretation, not speaker quotations.
Toolboxes: Viswajeeet Balaji (Microsoft), {VIDEO_URL}&t=4823s
Evolution of MCP auth: Den Delimarsky (Anthropic), {VIDEO_URL}&t=9149s
MCP Apps: Jeremiah Lowin (Prefect), {VIDEO_URL}&t=12065s
Event-driven agents: Clare Liguori (Amazon), {VIDEO_URL}&t=13508s

TRANSITION: "Before adding those capabilities, remove the shortcuts that already
prevent a prototype from being reliable."
""")


def release_checks(prs):
    slide = new_slide(prs, "Lessons from implementation | release tests still required", "Prototype shortcuts become production failures.")
    text(slide, "SHORTCUT", 0.65, 2.0, 5.4, 0.3, size=12, value_color=RED, bold=True)
    text(slide, "ENGINEERING RESPONSE", 6.6, 2.0, 6.0, 0.3, size=12, value_color=MSBLUE, bold=True)
    for index, (before, after) in enumerate([
        ("One environment for incompatible SDKs", "Isolate the MCP runtime; validate the boundary"),
        ("Unbounded tool processes", "Limit concurrency; timeout and reap children"),
        ("Shared mutable request context", "Create request-local workflow instances"),
        ("A plausible brief counts as success", "Evaluate factual coverage, refusal and recovery"),
    ]):
        vertical = 2.65 + index * 0.88
        text(slide, before, 0.65, vertical, 5.2, 0.68, size=21)
        arrow(slide, 5.92, vertical + 0.2, 0.38)
        text(slide, after, 6.6, vertical, 6.0, 0.7, size=21, value_color=MSBLUE)
        if index < 3:
            line(slide, 0.65, vertical + 0.74, 12.65, vertical + 0.74, LINE)
    text(slide, "Still to prove: unauthorized actor / malicious tool output / timeout replay / restart + replicas",
         0.65, 6.56, 12.0, 0.4, size=14, value_color=RED, bold=True)
    add_notes(slide, """
TIME: 19:30-20:30. Practical implementation lessons, then the close.

SAY: "The useful lessons were ordinary software engineering. Runtime dependencies
can conflict, so isolate the MCP environment and test its boundary. External tool
processes can hang or multiply, so bound concurrency, time out communication and
clean up children. Workflow state can leak across requests, so create request-local
instances. And a plausible answer can still omit the critical fact, so evaluate it."

POINT: Move across each row from shortcut to response. The current source has an
isolated bridge, a concurrency semaphore, timeout cleanup and a request-local workflow
wrapper. These are mechanisms present in code, not proof that all operational failure
modes have passed a fresh load or integration test. Avoid unsupported latency claims.

OPTIONAL DEPTH: Use golden cases for facts, refusal, unknown data and recovery. Check
the returned positions against totals, whether a focused answer stays focused, and
whether the supervisor preserves supplier uncertainty. A release decision should be
traceable to a version, configuration and evidence set. Keep a known-good version
available and rehearse rollback rather than treating a passing happy path as a gate.

CAUTION: The bottom line lists work still to prove: unauthorized actors, instructions
embedded in tool results, lost-response replay and restart/replica behavior. We are
not claiming those distributed guarantees or enterprise authorization exist in the
sample. Local idempotency is one useful result, not a production certificate.

TRANSITION: "That is the difference between a convincing demo and an accountable system."
HARD STOP: At 20:30 leave any unfinished demonstration and move to slide 14.
""")


def close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, dark=True, number=len(prs.slides))
    text(slide, "BUILD / SECURE / OBSERVE / SCALE", 0.65, 0.6, 11.9, 0.35,
         size=13, value_color=CYAN, bold=True)
    text(slide, "The model makes it capable.\nEngineering makes it accountable.",
         0.65, 1.65, 11.9, 2.1, size=40, font=DISPLAY, bold=True, value_color=WHITE)
    line(slide, 0.65, 4.3, 12.65, 4.3, CYAN)
    text(slide, "Discover typed tools. Inspect results. Test denied actions.",
         0.65, 4.7, 12.0, 0.55, size=25, value_color=CYAN, bold=True)
    text(slide, "Questions", 0.65, 6.04, 5.0, 0.5, size=25, value_color=CYAN, bold=True)
    text(slide, "Lee Stott | Global AI Community Bengaluru", 5.7, 6.15,
         6.95, 0.4, size=16, value_color=WHITE, align=PP_ALIGN.RIGHT)
    add_notes(slide, """
TIME: 20:30-22:00 close; 22:00-25:00 questions. Protect the final three minutes.

SAY: "MCP gives an application a consistent way to use tools. The workflow decides
how evidence moves. The application and domain controls decide who can act. Our
job is to make those responsibilities explicit, test the failure paths and retain
enough evidence to explain the outcome. The model makes the system capable.
Engineering makes it accountable."

RECAP: We located MCP in an actual architecture, chose sequential execution for
known dependencies, compared routing and parallel collaboration, designed a narrow
tool, demonstrated local human approval and replay, and examined hosted execution
and debugging. These are transferable patterns, not a claim that every production
requirement has been implemented in a synthetic sample.

QUESTION CUES: If asked why four agents, say the example makes evidence boundaries
visible; use fewer if the workload does not justify orchestration. If asked whether
the supervisor routes, say no: the graph is fixed and routing was a comparison.
If asked whether the local UI invokes Foundry, say no: the two demonstrations share
domain code and fixtures but have separate execution paths. If asked about approval
identity, distinguish the demo's typed name from enterprise authentication.

OPTIONAL FOLLOW-UP: Keep the UI ready to revisit denial or replay. Use the hidden
source appendix for timestamped MCP Live links and framework references. The detailed
backup decks are for technical discussion after the session, with old fixed claims
rechecked before use. Do not reopen setup or start a fresh deployment during questions.

TRANSITION: Invite questions about the engineering decisions, not just the tools.
""")


def sources(prs):
    slide = new_slide(prs, "Appendix | clickable sources", "Continue exploring MCP")
    resources = [
        ("MCP Community Connect Bengaluru | event and session", EVENT_URL),
        ("MCP architecture | clients, servers and tools", "https://modelcontextprotocol.io/docs/learn/architecture"),
        ("MCP tools | schemas, results and errors", "https://modelcontextprotocol.io/specification/2025-11-25/server/tools"),
        ("MCP lifecycle | initialization and capabilities", "https://modelcontextprotocol.io/specification/2025-11-25/basic/lifecycle"),
        ("Agent Framework | orchestration patterns", "https://github.com/microsoft/agent-framework/tree/main/python/packages/orchestrations"),
        ("MCP Live! | Toolboxes in Microsoft Foundry | 1:20:23", f"{VIDEO_URL}&t=4823s"),
        ("MCP Live! | Evolution of MCP auth | 2:32:29", f"{VIDEO_URL}&t=9149s"),
        ("MCP Live! | Building MCP Apps | 3:21:05", f"{VIDEO_URL}&t=12065s"),
        ("MCP Live! | Event-driven agents | 3:45:08", f"{VIDEO_URL}&t=13508s"),
        ("MCP Live! | official resource collection", "https://aka.ms/MCPLiveResources"),
    ]
    for index, (caption, url) in enumerate(resources):
        link_box = text(
            slide, caption, 0.75, 1.95 + index * 0.44, 11.9, 0.4,
            size=16, value_color=MSBLUE,
        )
        link_box.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    text(slide, "MCP Live additions are based on its published agenda, not a full transcript review.",
         0.75, 6.5, 11.9, 0.4, size=13, value_color=MUTED)
    slide._element.set("show", "0")
    add_notes(slide, """
APPENDIX: Hidden from the timed show. Use for follow-up and distribution.

PURPOSE: Give attendees the original material without turning a 25-minute session
into a reading list. Every resource label on this slide is clickable in PowerPoint.
The video links jump to the relevant published chapter starts. Encourage people to
watch the source sessions for implementation detail rather than treating our design
questions as a substitute for the original talks.

ATTRIBUTION: The event title, date and artwork are from the Global AI Community
event page. MCP Live topic selection is based on its published chapter list, not
a complete transcript review. The slide questions are original interpretations and
not quotations or endorsements by the speakers. The Agent Framework orchestration
reference was checked against current documentation during this deck revision.

CAUTION: Documentation evolves. Verify API versions, hosting capabilities and
authentication requirements when building, rather than copying a conference diagram
as deployment instructions. The diagram deliberately separates local deterministic
controls from hosted model execution. Neither the links nor the slide design imply
that unimplemented routing, parallel execution or enterprise controls were demonstrated.

PRESENTER ACTION: Keep this slide hidden during the main story. Navigate here manually
only if a question needs a source. For the live session use the notes and runbook;
for later study attendees can follow the links below. Sources accessed 14 September 2026.

""" + "\n".join(
        f"{caption}: {url}" for caption, url in resources
    ))


def mcp_primitives(prs):
    slide = new_slide(prs, "Technical appendix | roles + scope", "MCP is a protocol, not the agent's brain.")
    for horizontal, title_value, detail in [
        (0.65, "HOST", "Application + tool-use policy"),
        (4.8, "CLIENT", "Protocol session + requests"),
        (8.95, "SERVER", "Capabilities + tool handlers"),
    ]:
        block(slide, title_value, detail, horizontal, 2.05, 3.68)
    for horizontal in [4.39, 8.54]:
        arrow(slide, horizontal, 2.56, 0.3)
    for index, (name, purpose, status) in enumerate([
        ("Tools", "Callable operations", "DEMONSTRATED: inventory lookup"),
        ("Resources", "Addressable context / data", "NOT EXPOSED: recall document URI"),
        ("Prompts", "Reusable interaction templates", "NOT EXPOSED: assessment template"),
    ]):
        vertical = 3.67 + index * 0.84
        text(slide, name, 0.65, vertical, 2.0, 0.4, size=23, font=DISPLAY, bold=True)
        text(slide, purpose, 3.0, vertical, 4.0, 0.6, size=18)
        text(slide, status, 7.4, vertical, 5.2, 0.6, size=16, value_color=MSBLUE)
    text(slide, "REST can stay behind the server. MCP standardizes the AI-facing capability interface.",
         0.65, 6.6, 12.0, 0.4, size=17, value_color=RED, bold=True)
    slide._element.set("show", "0")
    add_notes(slide, """
TECHNICAL APPENDIX: Use after the timed talk, or answer a roles question in 60 seconds.

SAY: "The host is the application that controls the experience and tool-use policy.
It creates a client session to a server. The server exposes capabilities. A model can
propose a tool name and arguments, but application code mediates execution. A scripted
client can use the same server without any model; our terminal demonstration did exactly that."

WALKTHROUGH: Tools are callable operations such as locate_inventory. Resources are
addressable context, typically identified by a URI. Prompts are reusable templates
that a client can retrieve. The recall-document resource and assessment prompt on this
slide are possible extensions, not implemented endpoints. This server demonstrates tools
only. Do not call every returned JSON object an MCP resource or every system instruction
an MCP prompt. Those words also name specific protocol primitives.

WHY MCP INSTEAD OF JUST REST? The value is a common client-facing discovery and invocation
contract, not replacing every business API. A tool handler could call an existing REST
service, SQL query, or domain function. That integration code still exists. MCP may reduce
repeated adapters across compatible clients, but interoperability must be tested, including
capability support, authentication and error behavior. A single application may still be
well served by a direct function call; do not add a server without a reuse or boundary need.

TRANSPORT: Stdio uses a launched process's input/output streams; keep ordinary logs off
protocol stdout. Streamable HTTP is a remote transport with additional authentication and
network considerations. This sample's custom CLI/JSON bridge is neither a generic remote
MCP service nor an MCP stdio connection. Its MCP session runs in process inside the child.

CHECK UNDERSTANDING: Can the same MCP server serve a deterministic test client? Yes.
Does publishing a tool authorize its use? No. Does this sample expose resources or prompts?
No. These distinctions are more useful than memorizing a framework-specific class name.
""")


def mcp_code(prs):
    slide = new_slide(prs, "Technical appendix | repository code excerpts", "Follow one call from Python to MCP and back.")
    text(slide, "SERVER  /  caldova_mcp.py", 0.65, 2.0, 5.7, 0.38,
         size=19, font=DISPLAY, bold=True, value_color=MSBLUE)
    text(slide, '@mcp.tool(\n    title="Locate affected inventory",\n    annotations=ToolAnnotations(\n        read_only_hint=True,\n        open_world_hint=False),\n)\ndef locate_inventory(\n    batch_id: BatchId\n) -> dict[str, Any]:\n    return STORE.locate_inventory(batch_id)',
         0.65, 2.65, 5.95, 3.6, size=15, font="Cascadia Code")
    line(slide, 6.65, 2.0, 6.65, 6.3, LINE)
    text(slide, "CLIENT  /  inspect_mcp.py", 7.0, 2.0, 5.65, 0.38,
         size=19, font=DISPLAY, bold=True, value_color=MSBLUE)
    text(slide, 'async with Client(server) as client:\n    listing = await client.list_tools()\n    result = await client.call_tool(\n        "locate_inventory",\n        {"batch_id": "B-2408-AX7"},\n    )\n    facts = result.structured_content',
         7.0, 2.65, 5.65, 2.9, size=15, font="Cascadia Code")
    text(slide, "Validate the result and domain invariants.\nDo not scrape prose to recover inventory.",
         7.0, 5.55, 5.65, 0.8, size=18, value_color=MUTED)
    text(slide, "Imports, setup and docstring omitted. BatchId supplies the validated input constraints.",
         0.65, 6.62, 12.0, 0.32, size=14, value_color=MUTED)
    slide._element.set("show", "0")
    add_notes(slide, """
TECHNICAL APPENDIX: A line-by-line walkthrough for engineers. Excerpts are shortened
for legibility; use the repository files for runnable code. This SDK syntax matches
the installed MCP v2 environment, not all older examples found online.

SERVER: MCPServer registers capabilities. The decorator exposes one operation and its
metadata. ToolAnnotations reports hints, not enforced permissions. BatchId is an
Annotated string with Pydantic Field constraints: required input, length bounds and a
restricted pattern. The handler delegates to STORE, keeping protocol adaptation separate
from the business computation. That separation lets domain tests run without MCP and
protocol tests exercise the same logic through MCP.

OUTPUT: dict[str, Any] produces a generic object outputSchema in this environment. The
named fields seen in structured_content are actual data, not a strongly enforced DTO.
A production improvement would define a typed inventory result and validate invariants,
such as total_units matching the sum of positions. Do not present that improvement as
already implemented. The output should also be bounded so a large inventory cannot
consume an unbounded model context or response budget.

CLIENT: In inspect_mcp.py, server is StdioServerParameters pointing at caldova_mcp.py
using the current interpreter. Entering the client context starts the session and manages
initialization. list_tools returns metadata; call_tool invokes the named operation with
arguments. structured_content is the Python SDK spelling of wire-level structuredContent.
Leaving the context closes the session. The script adds a 30-second overall bound.

BUILD UP THE EXAMPLE: Start with discovery only. Add one valid inventory call. Finally
add invalid input, a schema-valid denied write, and a read to verify unchanged state.
That progression makes a protocol test explainable without a model. If the downstream
agent answer is wrong after these checks pass, inspect the handoff and synthesis next.

REPRODUCE FROM REPOSITORY ROOT:
./caldova-recall-control/.venv/Scripts/python.exe caldova-recall-control/scripts/inspect_mcp.py
Use .venv/bin/python instead on macOS/Linux. Install requirements-ui.txt first; do not
mix the Foundry hosting dependencies into this MCP v2 environment.
""")


def mcp_failures(prs):
    slide = new_slide(prs, "Technical appendix | test the boundary", "Different failures need different responses.")
    for index, (cause, observed, next_step) in enumerate([
        ("Malformed batch", "MCP isError: true", "Correct input; no blind retry"),
        ("Invalid approval credential", "MCP isError: true", "Require real authority"),
        ("Invalid UI approval handle", "HTTP 403 from FastAPI", "Application gate, not MCP RPC"),
        ("Correct facts, wrong brief", "Agent quality failure", "Check handoff + evaluation"),
    ]):
        vertical = 2.35 + index * 0.96
        text(slide, cause, 0.65, vertical, 4.0, 0.65, size=21, font=DISPLAY, bold=True)
        text(slide, observed, 4.9, vertical, 3.45, 0.65, size=18, value_color=MSBLUE)
        text(slide, next_step, 8.8, vertical, 3.85, 0.65, size=18)
        if index < 3:
            line(slide, 0.65, vertical + 0.77, 12.65, vertical + 0.77, LINE)
    text(slide, "Also test: unknown tool / timeout / malicious output / lost response / restart / replicas",
         0.65, 6.65, 12.0, 0.34, size=15, value_color=RED, bold=True)
    slide._element.set("show", "0")
    add_notes(slide, """
TECHNICAL APPENDIX: Use to make the demo's evidence precise and to discuss testing.

ROW 1: inspect_mcp.py sends 'bad batch!', which violates BatchId constraints. In the
installed SDK this returns a tool result with is_error true. The server has understood
the MCP request, but the tool cannot execute with these arguments. Correcting input is
different from retrying a transient outage. We display the observed error flag rather
than depending on an SDK-specific validation message in the live narration.

ROW 2: The script sends a synthetic 16-character hexadecimal credential. It passes the
input shape but has no matching approval in the domain store. The response is isError
true and sanitized text, 'Error executing tool quarantine_batch'. No usable credential
is created or printed. A second read verifies that the failed call left inventory
unchanged. This demonstrates policy enforcement, not verified human authentication.

ROW 3: The optional direct browser-API request supplies an invalid approval handle.
FastAPI rejects it with HTTP 403 before granting access to a valid approval credential.
That status is an application HTTP response, not an MCP JSON-RPC error. Separately,
unknown tools or protocol-level failures may use JSON-RPC error responses. Distinguish
the protocol's two error channels from the HTTP layer and from SDK exception settings.

ROW 4: A valid, factual tool response can still become an incomplete or misleading
brief. Domain tests cannot prove model quality. Evaluate factual coverage, uncertainty,
refusal to claim mutation, and focused-answer behavior against the actual tool evidence.

TEST LADDER: Run domain policy tests, then schema/annotation and in-process MCP tests,
then the stdio test, then API approval/replay tests, and finally hosted integration and
evaluation. A passing local suite is not proof that a cloud model, identity, trace, or
deployment is healthy. The tests live under caldova-recall-control/tests; the relevant
protocol file is test_caldova_mcp.py and the UI contract is test_control_tower_api.py.

REMAINING WORK: Exercise timeout and cancellation cleanup, tool-output prompt injection,
lost-response retries, restarts and replicas. These are follow-up tests or production
requirements, not guarantees demonstrated by this five-minute local sequence.
""")


ENGINEERING_NOTES = {
    "cover": """
LEARNING CONTRACT: By the end, an attendee should be able to locate the MCP client/server
boundary, inspect one discovered tool schema, call it without a model, interpret structured
results and failures, and distinguish protocol behavior from workflow and business policy.
For developers, the payoff is a repeatable integration and test pattern. For AI engineers,
it is separating tool evidence from model quality so failures can be evaluated accurately.
These are assessable outcomes, not a promise of complete production deployment training.

DELIVERY: The detailed notes are preparation and Q&A material, not a script to read
verbatim. Use SAY and the timed cues for the live talk; select one technical point per slide.
Ask: 'If the model confidently says stock is quarantined, what would you inspect to prove it?'
Return to that question when the result and audit are visible.
""",
    "scenario": """
ENGINEERING DETAIL: Turn the story into invariants: a known batch, exact stock exposure,
no unapproved mutation, and zero new mutations on replay. MCP transports the request and
result; the domain owns the inventory computation. The model does not get to invent the
stock count. This makes the example useful beyond pharmaceutical recall: the same split
appears in refund approval, order cancellation and infrastructure change workflows.

QUESTION: 'Could a single deterministic service do this?' Yes, and the local demo does.
The separate agent path explores unstructured requests and synthesis. Additional agents
should earn their complexity through measurable task decomposition, not through branding.
The saved UI's Model Router badge is not evidence that this browser path calls a model.
""",
    "architecture": """
SOURCE WALK: Open control_tower_api.py at invoke to show Client(mcp), then caldova_mcp.py
at locate_inventory to show its domain call. For the hosted lane, follow main.py's
locate_inventory wrapper into _call_mcp_read_tool and then mcp_v2_bridge.py call_tool.
Do not show credential-bearing environment files or deployment caches while navigating.

TRADEOFF: Isolating incompatible packages solves a dependency boundary, but adds process
startup overhead and another failure point. This sample does not provide a long-lived
remote MCP server, connection pool, shared state store, or dynamic discovery-driven agent
registration. Those would be explicit design choices, not consequences of using MCP.
CHECK: Ask which component decides the next specialist. Answer: the fixed workflow graph,
not the MCP server. Ask which enforces approval. Answer: application and domain policy.
""",
    "protocol": """
THREE CHECKS: Parse JSON; validate its shape; validate business meaning. These are separate
checks. The regex rejects malformed identifiers, but does not prove the batch exists.
An output schema can describe fields, but does not prove the inventory source is correct.
A readOnlyHint communicates intent, but does not constrain a malicious server's behavior.

API DESIGN: Prefer narrow names and bounded arguments over an execute_sql or run_command
tool. Describe unknown data explicitly. Avoid stuffing instructions into tool results;
the host should treat results as external evidence, not higher-priority instructions.
Contract changes need compatibility tests across consumers. Discovery helps clients learn
the shape, but does not guarantee that old prompts or client logic will interpret a new
field correctly. Check the hidden code appendix for the exact SDK-to-wire field mapping.
""",
    "patterns": """
DECISION RULE: Before adding agents, ask whether steps require separate instructions,
tools, evaluation criteria, or context. This demo's three read sources could be fetched
concurrently for the known batch; the implemented sequence is a teaching/design choice,
not proof that the data requires sequential access. A production comparison should
measure single-agent, deterministic fan-out, and multi-agent designs on quality, latency,
cost and failure recovery. MCP does not mandate any one orchestration pattern.

Q&A: 'Does the router in Model Router choose the next agent?' No. Model selection and
agent-workflow routing are different concerns. 'Is human approval an extra model vote?'
No. A separate authoritative policy check is required regardless of how evidence arrived.
""",
    "agents": """
CODE DETAIL: build_agents registers explicit Python function wrappers with each Agent.
Those wrappers call the read-only bridge. The application does not automatically give
the model everything returned by tools/list. Discovery and exposure are different steps.
The bridge's READ_ONLY_TOOLS allowlist provides a second check below the model-facing
tool list. The underlying demo MCP server still exposes seven tools to direct clients,
which is why a public remote deployment would need stronger server-side authorization.

EVALUATION: Check not only the supervisor's text, but each specialist's evidence and
handoff. Instructions such as 'call exactly once' are not enforced transaction guarantees.
The final response must not claim approval or mutation based on a persuasive request.
Keeping mutation tools out of the agent's tool set limits accidental execution but does
not make the whole server safe to expose to arbitrary clients.
""",
    "local_demo": """
MINUTE-BY-MINUTE: 6:00-7:15 terminal discovery/schema/call/failures; 7:15-8:15 browser reset
and analysis; 8:15-9:00 explain the optional API denial; 9:00-10:00 approve and act;
10:00-11:00 replay and inspect evidence. Skip the optional HTTP request if needed.
Do not install packages or start a new server while presenting. Rehearse the terminal's
scroll position so the seven numbered sections are easy to follow rather than reading
the full inventory JSON aloud. Keep the font large enough for the back of the room.

AUDIENCE PROMPT: Before section 6, ask: 'The credential passes the schema. Will the write
succeed?' After denial, ask which layer rejected it. Before replay, ask whether the total
stock count or positions_changed proves no duplicate mutation. The latter is the relevant
mutation evidence. This creates a predict-observe-explain sequence rather than a click tour.
""",
    "evidence": """
PROOF MATRIX: Terminal protocol/version/list/schema proves an MCP session and contract
inspection. The valid read proves the server returns synthetic structured facts. The
denied call plus unchanged read proves that invalid approval did not mutate that child's
state. The browser approval and replay prove the local application policy path. None of
these proves cloud deployment health, distributed transactions, or professional identity.

Q&A: 'Why does an idempotent operation return a different count on replay?' Because the
business state is unchanged while the response reports newly affected positions. 'What
if the process restarts?' The local store resets; do not generalize process-local behavior
to a durable multi-replica guarantee. A production design needs an atomic mutation and
idempotency record, with a defined key scope, retention period and approval lifecycle.
""",
    "debugging": """
DIAGNOSTIC ORDER: First inspect request arguments and schema. Then inspect transport or
process failure. Then inspect the server's structured result. Finally compare that result
with the agent handoff and final answer. Do not start by rewriting the system prompt when
the server is unavailable or has returned wrong data. Distinguish a tool result with
isError from an SDK exception, a JSON-RPC error, and the browser API's HTTP response.

METRICS: Track counts and latency separately for model calls, MCP tool calls, and custom
bridge startup. One slow request is not a benchmark. Record tool name, version, sanitized
arguments, outcome and correlation where supported. Trace propagation across the custom
subprocess boundary is not automatic. Demonstrate only spans actually present and never
put reusable approval credentials into logs merely to make a trace look comprehensive.
""",
    "hosted_demo": """
WHAT THIS ADDS: The local client already proved the tool contract without a model. Now
the model-backed workflow must select its allowed tool, preserve evidence through the
fixed graph, and produce a grounded brief. Evaluate coverage and refusal independently
from protocol success. A successful tools/call does not imply that the agent used its
result faithfully. Conversely, a poor summary does not prove the MCP server failed.

CODE FALLBACK: Show build_agents tool lists and build_workflow edges if hosted access is
unavailable. Be explicit that source shows intended configuration, not executed behavior.
The bridge returns structured data from its in-process MCP session through custom JSON
IPC. The new terminal demo shows real stdio separately. Never describe the custom bridge
as a generic MCP gateway or claim discovery automatically configured these agents.
""",
    "production": """
THREAT MODEL: Tool descriptions and results may contain hostile instructions. A caller
may bypass the UI, replay an old handle, ask the model to invent approval, or target a
different batch. Protect the exact principal/resource/action combination and treat model
output as a recommendation, not an authorization token. A prompt and an annotation are
not enforceable access controls. An authenticated workload is not automatically entitled
to every action exposed by its downstream service.

REMOTE MCP: Stdio reachability and remote HTTP authentication are different concerns.
For a remote service, follow the applicable MCP authorization specification and validate
the intended audience and resource. Do not forward unrelated access tokens to arbitrary
tools. This sample does not demonstrate that remote authorization flow; the slide names
requirements to design and validate, not configuration already deployed by the demo.
""",
    "mcp_live": """
CONNECT TO THE DEMO: Tool curation changes the discovered/exposed capability set; inspect
what each workload can see. Authentication adds caller identity; it does not decide the
pharmacist's business authority by itself. An embedded interface may clarify evidence;
it must not smuggle approval around server policy. An event can start analysis; duplicate
delivery still needs deduplication and should not silently trigger a stock mutation.

AUDIENCE EXERCISE: Pick one extension and name its new failure case before choosing a
library. For event-driven analysis, what happens when the recall event is delivered twice?
For a new UI, can a caller invoke the endpoint directly? Keep the discussion in terms of
testable boundaries. Do not imply that the published MCP Live agenda proves any of these
integrations exists in this repository or has been tested with this version of the SDK.
""",
    "release_checks": """
REPRODUCIBLE CHECKS: From the repository root run the MCP inspection script, then
./caldova-recall-control/.venv/Scripts/python.exe -m pytest caldova-recall-control/tests -q
The protocol suite checks discovery, schemas, annotations, structured facts, stdio,
sanitized denial and replay. The API suite checks the separate browser-facing policy.
Local tests do not contact the deployed model to certify its live answers. Publish a
fresh versioned evidence set instead of repeating a historical count from a slide.

CHANGE REVIEW: If a tool becomes destructive, its name, description, annotations, schema,
allowlists, authorization and tests all deserve review. If a dependency changes, rerun
both environments and their integration boundary. If the model changes, rerun evaluation.
Source compatibility, protocol compatibility, policy correctness and model quality are
different gates. The hidden failure appendix gives a compact checklist for discussion.
""",
    "close": """
KNOWLEDGE CHECK: Ask three questions. What method discovers tool metadata? tools/list.
What does structuredContent prove? It provides machine-readable evidence, not guaranteed
truth. What stops the denied write? Application/domain policy, not readOnlyHint or the
model's intentions. Ask a fourth if time permits: did the hosted bridge use MCP stdio?
No; the dedicated terminal demonstration did, while the bridge used in-process MCP.

TAKE-HOME TASK: Run inspect_mcp.py and inspect BatchId. Then propose, without silently
changing the live contract, a typed inventory output and an invariant test for the sum
of positions. That exercise links protocol discovery, schema design and domain testing.
Use slides 16-18 for follow-up questions on roles, code and failures. The detailed notes
companion can be read without PowerPoint; the main talk remains 14 slides plus questions.
""",
}


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = SESSION_TITLE
    prs.core_properties.subject = "MCP Community Connect Bengaluru | 26 September 2026"
    prs.core_properties.author = "Lee Stott"
    companion = [
        "# MCP Community Connect Bengaluru: Speaker Notes",
        "14 main slides for a 25-minute session; four hidden reference/deep-dive slides. "
        "SAY and TIME are the live talk track. Engineering detail and appendices are "
        "preparation and Q&A material, not text to read verbatim.",
        "Learning outcomes: identify the MCP boundary; inspect discovery and schemas; "
        "call tools without a model; distinguish structured evidence, errors, orchestration, "
        "and business authorization. All recall data is synthetic.",
    ]
    for builder in [
        cover, scenario, architecture, protocol, patterns, agents, local_demo,
        evidence, debugging, hosted_demo, production, mcp_live, release_checks,
        close, sources, mcp_primitives, mcp_code, mcp_failures,
    ]:
        builder(prs)
        slide = prs.slides[-1]
        notes = slide.notes_slide.notes_text_frame.text
        if builder.__name__ in ENGINEERING_NOTES:
            notes += "\n\nOPTIONAL ENGINEERING DETAIL / Q&A\n" + ENGINEERING_NOTES[builder.__name__].strip()
            add_notes(slide, notes)
        companion.append(f"## Slide {len(prs.slides)}: {builder.__name__.replace('_', ' ').title()}\n\n{notes}")
    output = ROOT / "mcp-community-connect-bengaluru-speaker.pptx"
    prs.save(output)
    (ROOT / "mcp-community-connect-bengaluru-speaker-notes.md").write_text(
        "\n\n".join(companion) + "\n", encoding="utf-8"
    )
    return output


if __name__ == "__main__":
    print(f"wrote {build()}")