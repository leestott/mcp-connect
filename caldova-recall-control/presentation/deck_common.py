"""Shared MCP Community Connect Bengaluru theme and slide primitives."""

from __future__ import annotations

from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Microsoft / Fluent palette
INK = "182B35"          # primary text on light
PAPER = "F7FAFA"        # neutral light background
WHITE = "FFFFFF"
MUTED = "50656B"        # secondary text
LINE = "D6E3E5"         # neutral stroke
MSBLUE = "176D80"
MSBLUE_DK = "164B68"
DEEP = "102D3C"
GREEN = "107C10"
RED = "BC4434"
AMBER = "FFB900"
PURPLE = "5C2D91"
CYAN = "78D3D0"

# tints
BLUE_SOFT = "E1F1F1"
GREEN_SOFT = "DFF6DD"
RED_SOFT = "FDE7E9"
AMBER_SOFT = "FFF4CE"

EVENT_NAME = "MCP COMMUNITY CONNECT | BENGALURU"
EVENT_DATE = "26 SEPTEMBER 2026"
EVENT_ART = Path(__file__).resolve().parent / "assets" / "mcp-community-connect-bengaluru.png"

# Typography — Microsoft brand fonts
DISPLAY = "Segoe UI Semibold"
BODY = "Segoe UI"
LIGHT = "Segoe UI Light"
MONO = "Cascadia Code"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def fill(shape, value: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(value)
    shape.line.fill.background()


def rect(slide, x, y, w, h, value, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, value)
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1)
    return shape


def text(slide, value, x, y, w, h, *, size=18, value_color=INK, font=BODY,
         bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color(value_color)
    return box


def line(slide, x1, y1, x2, y2, value=LINE, width=1.0, dash=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(value)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def logo_mark(slide, x, y, unit=0.15, gap=0.03):
    """Place the official event artwork without cropping or recoloring it."""
    return slide.shapes.add_picture(str(EVENT_ART), Inches(x), Inches(y),
                                    width=Inches(unit * 2 + gap))


def bullet_row(slide, value, x, y, w, *, size=14, value_color=INK, accent=MSBLUE,
               font=BODY, bold=False, height=0.5):
    rect(slide, x, y + 0.05, 0.12, 0.12, accent)
    text(slide, value, x + 0.32, y - 0.03, w - 0.32, height, size=size,
         value_color=value_color, font=font, bold=bold)


def code_line(slide, value, x, y, w, *, size=11, value_color="D6E4F0", height=0.26):
    text(slide, value, x, y, w, height, size=size, value_color=value_color, font=MONO)


def add_notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value.strip()


def base(slide, dark=False, number=None, brand="CALDOVA RECALL CONTROL TOWER"):
    rect(slide, 0, 0, 13.333, 7.5, DEEP if dark else PAPER)
    foreground = "B6CFD4" if dark else MUTED
    line(slide, 0.55, 7.12, 12.78, 7.12, MSBLUE if dark else LINE)
    text(slide, EVENT_NAME, 0.55, 7.22, 4.45, 0.18, size=8,
        value_color=foreground, bold=True)
    text(slide, brand, 5.05, 7.22, 4.65, 0.18, size=7.5, value_color=foreground)
    text(slide, EVENT_DATE, 10.0, 7.22, 2.1, 0.18, size=8, value_color=foreground)
    if number is not None:
       text(slide, f"{number:02d}", 12.42, 7.2, 0.36, 0.2, size=9,
           value_color=foreground, bold=True, align=PP_ALIGN.RIGHT)


def heading(slide, kicker, title_value, number):
    base(slide, number=number)
    text(slide, kicker.upper(), 0.55, 0.42, 8.0, 0.22, size=9, value_color=MSBLUE, bold=True)
    text(slide, title_value, 0.55, 0.68, 12.2, 0.62, size=30, font=DISPLAY, bold=True)
    line(slide, 0.55, 1.4, 12.78, 1.4, LINE)


def label(slide, value, x, y, w, background, foreground=WHITE):
    rect(slide, x, y, w, 0.28, background)
    text(slide, value.upper(), x + 0.08, y + 0.05, w - 0.16, 0.14, size=8,
         value_color=foreground, bold=True)


def node(slide, title_value, subtitle, x, y, w, h, accent=MSBLUE):
    rect(slide, x, y, w, h, WHITE, line=LINE)
    rect(slide, x, y, 0.07, h, accent)
    text(slide, title_value, x + 0.18, y + 0.15, w - 0.3, 0.28, size=14, font=DISPLAY, bold=True)
    text(slide, subtitle, x + 0.18, y + 0.48, w - 0.3, h - 0.58, size=9.5, value_color=MUTED)


def title_slide(prs, *, demo_tag, title_value, subtitle, chips, disclaimer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide)
    rect(slide, 0, 0, 13.333, 0.12, MSBLUE)
    text(slide, "MCP Community Connect", 0.65, 0.55, 8.0, 0.5,
        size=26, font=DISPLAY, bold=True)
    text(slide, "BENGALURU  /  GLOBAL AI COMMUNITY", 0.68, 1.12, 8.0, 0.25,
        size=11, value_color=MSBLUE, bold=True)
    text(slide, demo_tag.upper(), 0.68, 1.8, 7.3, 0.28,
        size=11, value_color=MSBLUE, bold=True)
    text(slide, title_value, 0.65, 2.27, 7.2, 1.55, size=38,
        font=DISPLAY, bold=True)
    text(slide, subtitle, 0.68, 4.0, 7.0, 1.05, size=17, value_color=MUTED)
    slide.shapes.add_picture(str(EVENT_ART), Inches(8.15), Inches(1.48), width=Inches(4.5))
    text(slide, "Lee Stott", 0.68, 5.32, 3.2, 0.32, size=17, font=DISPLAY, bold=True)
    text(slide, "Principal Cloud Advocate Manager, Microsoft", 0.68, 5.72, 7.0, 0.25,
        size=11, value_color=MUTED)
    text(slide, "26 September 2026 | Samarthanam Auditorium", 8.05, 6.02, 4.7, 0.25,
        size=10, value_color=MSBLUE, align=PP_ALIGN.CENTER)
    for index, chip in enumerate(chips):
       x = 0.68 + index * 3.05
       rect(slide, x, 6.42, 0.05, 0.26, MSBLUE)
       text(slide, chip, x + 0.14, 6.42, 2.82, 0.3, size=10,
           value_color=MSBLUE_DK, bold=True)
    text(slide, disclaimer, 0.68, 6.87, 11.8, 0.18, size=8, value_color=MUTED)
    return slide


def section_slide(prs, *, index_value, kicker, title_value, detail):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, dark=True)
    text(slide, index_value, 0.7, 0.85, 3.0, 1.2, size=64, value_color="24405E", font=DISPLAY, bold=True)
    text(slide, kicker.upper(), 0.75, 2.5, 8.0, 0.3, size=11, value_color=CYAN, bold=True)
    text(slide, title_value, 0.72, 2.95, 11.5, 1.6, size=38, value_color=WHITE, font=DISPLAY, bold=True)
    text(slide, detail, 0.75, 4.7, 10.8, 1.2, size=17, value_color="C7D6E6")
    text(slide, "BUILD / SECURE / OBSERVE / SCALE / CONNECT", 0.75, 6.62, 8.0, 0.25,
         size=10, value_color=CYAN, bold=True)
    return slide


def demo_slide(prs, *, number, demo_name, goal, steps, show, fallback, command=None):
    """A prominent, unmistakable live-demo placeholder slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, dark=True)
    rect(slide, 0, 0, 13.333, 1.5, MSBLUE)
    text(slide, "LIVE DEMO", 0.7, 0.34, 5.0, 0.4, size=13, value_color="BFE6FF", bold=True)
    text(slide, demo_name, 0.7, 0.66, 11.0, 0.7, size=30, value_color=WHITE, font=DISPLAY, bold=True)
    text(slide, f"{number:02d}", 12.3, 0.4, 0.6, 0.6, size=26, value_color="BFE6FF", font=DISPLAY, bold=True, align=PP_ALIGN.RIGHT)

    text(slide, "GOAL", 0.7, 1.75, 5.0, 0.22, size=9, value_color=CYAN, bold=True)
    text(slide, goal, 0.7, 2.02, 11.8, 0.6, size=16, value_color="E4ECF5")

    text(slide, "RUN", 0.7, 2.85, 5.0, 0.22, size=9, value_color=CYAN, bold=True)
    for index, step in enumerate(steps):
        y = 3.15 + index * 0.5
        text(slide, f"{index + 1:02d}", 0.7, y, 0.45, 0.3, size=11, value_color=MSBLUE, font=MONO, bold=True)
        text(slide, step, 1.2, y - 0.02, 6.0, 0.45, size=12.5, value_color="E4ECF5")

    rect(slide, 7.5, 2.85, 5.15, 3.3, "122A47")
    text(slide, "SHOW ON SCREEN", 7.75, 3.05, 4.6, 0.22, size=9, value_color=CYAN, bold=True)
    for index, item in enumerate(show):
        y = 3.4 + index * 0.48
        rect(slide, 7.78, y + 0.05, 0.12, 0.12, GREEN)
        text(slide, item, 8.05, y - 0.03, 4.45, 0.45, size=11.5, value_color="E4ECF5")

    if command:
        rect(slide, 0.7, 6.35, 6.5, 0.5, "0A1A30")
        code_line(slide, command, 0.9, 6.47, 6.2, size=10.5, value_color=CYAN)
    rect(slide, 7.5, 6.35, 5.15, 0.5, "2A1B00")
    text(slide, f"Fallback: {fallback}", 7.7, 6.46, 4.8, 0.3, size=9.5, value_color=AMBER)
    return slide


def closing_slide(prs, *, title_value, rules, proof_headline, proof_caption, proof_items, resources):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, dark=True)
    text(slide, "GLOBAL AI COMMUNITY | BENGALURU", 0.62, 0.63, 6.4, 0.3,
         size=11, value_color=CYAN, bold=True)
    text(slide, title_value, 0.62, 1.2, 6.6, 1.5, size=34, value_color=WHITE, font=DISPLAY, bold=True)
    for index, (num, rule) in enumerate(rules):
        x = 0.65 + (index % 2) * 3.3
        y = 3.05 + (index // 2) * 1.2
        text(slide, num, x, y, 0.42, 0.22, size=11, value_color=CYAN, bold=True)
        text(slide, rule, x + 0.5, y - 0.05, 2.6, 0.7, size=16, value_color=WHITE, font=DISPLAY, bold=True)

    rect(slide, 7.5, 0, 5.83, 7.5, PAPER)
    text(slide, "PROOF", 7.9, 0.78, 2.1, 0.22, size=9, value_color=MSBLUE, bold=True)
    text(slide, proof_headline, 7.9, 1.3, 5.0, 0.78, size=50, value_color=RED, font=DISPLAY, bold=True)
    text(slide, proof_caption, 7.93, 2.16, 5.0, 0.32, size=17, font=DISPLAY, bold=True)
    line(slide, 7.93, 2.78, 12.9, 2.78, LINE)
    for index, item in enumerate(proof_items):
        rect(slide, 7.95, 3.2 + index * 0.52, 0.12, 0.12, MSBLUE)
        text(slide, item, 8.22, 3.12 + index * 0.52, 4.5, 0.3, size=13, value_color=INK)
    line(slide, 7.93, 6.05, 12.9, 6.05, LINE)
    text(slide, "Learn more", 7.9, 6.2, 3.0, 0.3, size=11, value_color=MSBLUE, bold=True)
    text(slide, resources, 7.9, 6.5, 5.0, 0.6, size=10, value_color=MUTED)
    return slide
