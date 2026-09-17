"""Build the story-led event edition without changing the original deck."""

import re
from pathlib import Path

from pptx import Presentation

import build_event_deck as theme
from deck_common import SLIDE_H, SLIDE_W, add_notes


ROOT = Path(__file__).resolve().parent
SCRIPT = ROOT / "mcp-community-connect-bengaluru-story-stage-script.md"
OUTPUT = ROOT / "mcp-community-connect-bengaluru-final.pptx"
STORY_OUTPUT = ROOT / "mcp-community-connect-bengaluru-story.pptx"
SLIDES = [
    ("cover", "The incident", "Asha needs a decision. Not another answer.",
     "From Prototype to Production", []),
    ("scenario", "Problem | 0:45-2:00", "One alert. Four locations. A person accountable.",
     "Define success before choosing an agent.",
     [("2,196", "synthetic units affected"), ("4 locations", "one batch to assess"),
      ("Asha Rao", "must decide whether to approve")]),
    ("flow", "Decision | 2:00-3:00", "Give the application a common tool contract.",
     "MCP standardises tool access. Your application still owns policy.",
     [("HOST + CLIENT", "Own the session and tool-use policy"),
      ("MCP SERVER", "Discover tools; validate calls; return evidence"),
      ("DOMAIN / APIs", "Inventory facts and business rules")]),
    ("contract", "Demonstration | 3:00-5:00", "Find the stock before asking for a decision.",
     "A narrow tool turns a vague request into checkable evidence.",
     [("DISCOVER", "tools/list: inspect locate_inventory"),
      ("CALL", "tools/call: supply the known batch_id"),
      ("CHECK", "Expected: 2,196 units across 4 positions")]),
    ("rows", "Decision | 5:00-6:30", "Which workflow serves Asha's question?",
     "Choose by dependencies and uncertainty, not by agent count.",
     [("SEQUENTIAL | demonstrated", "Build a brief through explicit specialist handoffs."),
      ("SUPERVISOR ROUTING | comparison", "Choose a specialist for a variable request; bound loops."),
      ("PARALLEL | comparison", "Fetch independent evidence; define partial-failure policy.")]),
    ("flow", "Demonstration | 6:30-9:00", "Turn tool evidence into a decision brief.",
    "Run analysis in the hosted Control Tower; retain the response ID.",
     [("TRIAGE", "Establish the batch"), ("INVENTORY", "Measure exposure"),
      ("COMPLIANCE", "Preserve uncertainty"), ("SUPERVISOR", "Synthesise; no tools")]),
    ("rows", "Failure lab | 9:00-13:00", "The request failed. What should happen next?",
     "Fix the layer that owns the failure. Verify the resulting state.",
     [("BAD INPUT | inspect_mcp.py", "Observe isError; correct the argument, not the prompt."),
      ("NO APPROVAL | inspect_mcp.py", "Observe denial; re-read stock to check no mutation."),
      ("WRONG BRIEF | diagnostic comparison", "Compare tool result, handoff and synthesis in that order.")]),
    ("demo", "Demonstration | 13:00-16:00", "Asha can approve. The model cannot.",
     "Human approval is a server-enforced policy boundary.",
    [("ASSESS", "Retained Foundry response; read-only agent"),
     ("APPROVE", "Authenticated person authorises quarantine"),
      ("INSPECT", "Check the operation result and audit")]),
    ("evidence", "Evidence | 16:00-17:00", "A retry must not become a second action.",
     "Measure the business effect, not the confidence of the answer.",
     [("FIRST AUTHORISED CALL", "Expected: 4 positions changed; 2,196 units processed"),
      ("REPLAY", "Expected: 0 additional positions changed"),
      ("AUDIT", "Decision and outcome; no reusable approval token")]),
    ("rows", "Production decisions | 17:00-20:00", "Now put real users and replicas around it.",
     "Hosting is a step. Identity, durable state and release evidence are gates.",
     [("WHO CAN ACT?", "Verified caller; scoped approval; least-privilege workload."),
      ("WHAT SURVIVES A RETRY?", "Atomic mutation + idempotency record; defined expiry."),
      ("WHAT SUPPORTS RELEASE?", "Traces, evaluations, reviewed tool changes and rollback.")]),
    ("rows", "Lesson | 20:00-21:00", "The same decisions apply to your workload.",
     "MCP makes tools accessible. Engineering makes actions accountable.",
     [("CHOOSE + DESIGN", "Smallest useful workflow; narrow, bounded tool contracts."),
      ("INVESTIGATE", "Find the first divergence; separate tool and model quality."),
      ("CONTROL", "Enforce authority; test denial, retries and state changes.")]),
    ("close", "Outcome | 21:00-22:00 | Questions 22:00-25:00", "Would you authorise the action?",
     "Choose a workflow. Design its tools. Investigate failures. Control actions.",
     [("FACTS", "Where did the answer come from?"),
      ("AUTHORITY", "Who is allowed to approve?"),
      ("EVIDENCE", "What changed, and what did the retry do?")]),
]
MAIN_SEQUENCE = [
    1, 2, 3, "architecture", "demo-1", 4, 5, "demo-2", 6,
    "demo-3", 7, "demo-4", 8, "demo-5", 9, 10, 11, 12,
]
DEMO_BREAKS = {
    "demo-1": ("01", "Find the stock", "MCP terminal", "Discover the tool, inspect its contract, and verify the inventory facts."),
    "demo-2": ("02", "Produce the brief", "Hosted Control Tower", "Run the read-only agent workflow and retain its response evidence."),
    "demo-3": ("03", "Investigate a failure", "MCP terminal", "Separate invalid input from denied authority, then inspect unchanged state."),
    "demo-4": ("04", "Give authority to the person", "Hosted Control Tower", "Move from recommendation to an authenticated, application-controlled action."),
    "demo-5": ("05", "Prove the outcome", "Hosted Control Tower", "Replay the operation and compare the state change with the audit evidence."),
}


def demo_break(prs, number, title, destination, purpose):
    slide = theme.new_slide(prs, f"DEMO {number} | SWITCH TO {destination.upper()}", title)
    theme.rect(slide, 0.65, 2.2, 2.15, 2.15, theme.MSBLUE)
    theme.text(slide, "DEMO", 0.65, 2.65, 2.15, 0.35, size=20, bold=True,
               value_color=theme.WHITE, align=theme.PP_ALIGN.CENTER)
    theme.text(slide, number, 0.65, 3.05, 2.15, 0.65, size=42, bold=True,
               value_color=theme.WHITE, align=theme.PP_ALIGN.CENTER)
    theme.text(slide, purpose, 3.35, 2.42, 8.8, 1.25, size=28)
    theme.text(slide, f"NEXT: {destination}", 3.35, 4.0, 8.8, 0.4, size=16,
               bold=True, value_color=theme.MSBLUE)
    add_notes(slide, f"DEMO BREAK {number}: {title}.\n\n"
              f"ACTION: Switch to the prepared {destination}.\n"
              f"PURPOSE: {purpose}\n"
              "Do not advance until the demonstration surface is ready. Use the scripted fallback if live evidence is unavailable.")
    return slide


def story_architecture(prs):
    slide = theme.new_slide(prs, "Application architecture | before the demonstrations",
                            "One application. Separate reasoning and authority paths.")
    theme.text(slide, "READ-ONLY ASSESSMENT", 0.65, 1.72, 5.8, 0.3, size=13,
               bold=True, value_color=theme.MSBLUE)
    for horizontal, title, detail in [
        (0.65, "Control Tower", "Browser"),
        (3.7, "FastAPI host", "Hosted analysis"),
        (6.75, "Foundry agent", "Agent Framework"),
        (9.8, "MCP read tools", "Synthetic fixtures"),
    ]:
        theme.block(slide, title, detail, horizontal, 2.15, 2.55)
    for horizontal in [3.25, 6.3, 9.35]:
        theme.arrow(slide, horizontal, 2.68, 0.34)
    theme.text(slide, "AUTHORISED ACTION", 0.65, 3.62, 5.8, 0.3, size=13,
               bold=True, value_color=theme.RED)
    for horizontal, title, detail in [
        (0.65, "EasyAuth identity", "Verified caller"),
        (3.7, "Application policy", "Bound approval"),
        (6.75, "Domain action", "Quarantine + audit"),
        (9.8, "Blob session", "ETag + replay state"),
    ]:
        theme.block(slide, title, detail, horizontal, 4.05, 2.55)
    for horizontal in [3.25, 6.3, 9.35]:
        theme.arrow(slide, horizontal, 4.58, 0.34)
    theme.text(slide, "SEPARATE MCP TERMINAL DEMO", 0.65, 5.57, 3.0, 0.28, size=12,
               bold=True, value_color=theme.MSBLUE)
    theme.text(slide, "Inspector -> MCP stdio server -> isolated synthetic state",
               3.65, 5.54, 8.7, 0.35, size=18)
    theme.text(slide, "The hosted agent can recommend. Only application policy can authorise and mutate.",
               0.65, 6.38, 12.0, 0.42, size=17, bold=True, value_color=theme.RED)
    add_notes(slide, """TIME: Immediately before Demo 1. Introduce the application architecture.

SAY: "The Control Tower has two deliberately different paths. Across the top, the browser asks the FastAPI host for an assessment. The Foundry-hosted Agent Framework workflow can use only read tools over the application's isolated MCP bridge. It can prepare a recommendation; it cannot approve or quarantine inventory.

Across the second lane, EasyAuth supplies the caller identity. Application policy binds approval to that caller, batch, action and session generation. The domain action records the quarantine and audit result, while the actor-scoped Blob session uses optimistic concurrency and retains replay state.

Our terminal demonstration is separate again: the inspector talks genuine MCP stdio to an isolated synthetic server. Its state is not the browser session."

POINT: Follow each lane left to right. Do not imply that the model grants authority or that the terminal and hosted browser share live state.
TRANSITION: "With those boundaries visible, let's inspect the first tool without a model."""
    )
    return slide


def draw(slide, layout, rows):
    if layout in {"scenario", "demo"}:
        for index, (label, detail) in enumerate(rows):
            vertical = 2.0 + index * 1.35
            theme.text(slide, label, 0.65, vertical, 3.6, 0.6,
                       size=34 if layout == "scenario" else 24, bold=True,
                       value_color=theme.RED if index == 0 else theme.MSBLUE)
            theme.text(slide, detail, 0.65, vertical + 0.65, 3.6, 0.65, size=17)
        theme.picture_fit(slide, "control-tower-initial.png", 4.65, 2.0, 8.0, 4.1)
        theme.text(slide, "Saved local UI | synthetic data | live results must be checked",
                   4.65, 6.05, 8.0, 0.4, size=12, value_color=theme.MUTED)
    elif layout == "flow":
        width = (12.0 - 0.35 * (len(rows) - 1)) / len(rows)
        for index, (label, detail) in enumerate(rows):
            horizontal = 0.65 + index * (width + 0.35)
            theme.text(slide, label, horizontal, 2.65, width, 0.5,
                       size=21, bold=True, value_color=theme.MSBLUE)
            theme.line(slide, horizontal, 3.35, horizontal + width, 3.35, theme.MSBLUE, 3)
            theme.text(slide, detail, horizontal, 3.7, width, 1.1, size=21)
            if index < len(rows) - 1:
                theme.arrow(slide, horizontal + width + 0.04, 3.35, 0.25)
        detail = ("Microsoft Agent Framework: fixed graph | Microsoft Foundry: hosted execution"
                  if len(rows) == 4 else "Discover -> call -> structured result | no model required for this check")
        theme.text(slide, detail, 0.65, 5.55, 12, 0.65, size=20, value_color=theme.MUTED)
    elif layout == "contract":
        for index, (label, detail) in enumerate(rows):
            vertical = 2.1 + index * 1.25
            theme.text(slide, label, 0.65, vertical, 2.0, 0.5, size=22,
                       bold=True, value_color=theme.MSBLUE)
            theme.text(slide, detail, 2.8, vertical, 9.5, 0.7, size=24)
        theme.text(slide, "Required batch_id | 1-64 characters | uppercase letters, digits, hyphens",
                   0.65, 5.95, 12, 0.45, size=17, value_color=theme.MUTED)
    else:
        for index, (label, detail) in enumerate(rows):
            vertical = 2.05 + index * 1.4
            theme.text(slide, label, 0.65, vertical, 11.9, 0.45, size=18,
                       bold=True, value_color=theme.MSBLUE)
            theme.text(slide, detail, 0.65, vertical + 0.53, 11.9, 0.65,
                       size=27 if layout in {"close", "evidence"} else 24)
            theme.line(slide, 0.65, vertical + 1.2, 12.6, vertical + 1.2, theme.LINE)


def build():
    script = SCRIPT.read_text(encoding="utf-8")
    sections = re.findall(r"^## Slide (\d+):[^\n]*\n(.*?)(?=^## |\Z)", script, re.M | re.S)
    if [int(number) for number, _ in sections] != list(range(1, 13)):
        raise ValueError("Stage script must contain slides 1 through 12 in order")
    section_notes = {int(number): notes for number, notes in sections}
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = theme.SESSION_TITLE + " | Story Edition"
    prs.core_properties.author = "Lee Stott"
    prs.core_properties.subject = "Synthetic recall: problem, decision, demonstration, evidence, lesson"
    demo_slides = {}
    for item in MAIN_SEQUENCE:
        if item == "architecture":
            story_architecture(prs)
        elif isinstance(item, str):
            demo_slides[item] = demo_break(prs, *DEMO_BREAKS[item])
        else:
            layout, kicker, title, lesson, rows = SLIDES[item - 1]
            if layout == "cover":
                theme.cover(prs)
            else:
                slide = theme.new_slide(prs, kicker, title)
                draw(slide, layout, rows)
                theme.text(slide, lesson, 0.65, 6.56, 12.0, 0.42,
                           size=16, value_color=theme.MSBLUE, bold=True)
            add_notes(prs.slides[-1], section_notes[item])
    for builder in [theme.architecture, theme.mcp_primitives, theme.mcp_code,
                    theme.mcp_failures, theme.mcp_live, theme.sources]:
        builder(prs)
        slide = prs.slides[-1]
        slide._element.set("show", "0")
        notes = slide.notes_slide.notes_text_frame.text
        add_notes(slide, "STORY EDITION APPENDIX: outside the timed show.\n"
                  "Original-edition timing and navigation cues below do not apply.\n\n" + notes)
    demo_index = theme.new_slide(prs, "Presenter navigation | outside timed show", "Demo Index")
    demo_index._element.set("show", "0")
    destinations = [
        ("demo-1", "Find the Stock", "Prepared MCP terminal: discovery and valid read"),
        ("demo-2", "Produce the Brief", "Hosted Control Tower: Run analysis; inspect response ID"),
        ("demo-3", "Investigate a Failure", "Same MCP terminal: validation, denial, unchanged state"),
        ("demo-4", "Give Authority to the Person", "Same hosted session: authenticated approval and audit"),
        ("demo-5", "Prove the Outcome", "Keep browser open: first call and latest replay"),
    ]
    for index, (key, title, destination) in enumerate(destinations):
        vertical = 2.0 + index * 0.88
        number = int(key[-1])
        link = theme.text(demo_index, f"{number:02d}  {title}", 0.65, vertical, 5.6, 0.5,
                          size=22, bold=True, value_color=theme.MSBLUE)
        link.click_action.target_slide = demo_slides[key]
        theme.text(demo_index, destination, 6.4, vertical, 6.0, 0.65, size=18)
    add_notes(demo_index, "Hidden navigation slide, outside the 25-minute schedule. "
              "Select a title to return to its main slide. Browser links do not start servers. "
              "Switch to prepared terminal/hosted windows manually; no executable links. "
              "Use only genuine saved evidence or explicitly labelled source fallbacks.")
    for number, slide in enumerate(list(prs.slides)[:18], start=1):
        if number == 1:
            continue
        link = theme.text(slide, "Demo Index", 11.1, 0.4, 1.55, 0.3,
                          size=11, bold=True, value_color=theme.MSBLUE)
        link.click_action.target_slide = demo_index
    prs.save(OUTPUT)
    prs.save(STORY_OUTPUT)
    check = Presentation(OUTPUT)
    assert len(check.slides) == 25
    assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in check.slides)
    assert all(slide._element.get("show") == "0" for slide in list(check.slides)[18:])
    for number in range(2, 19):
        link = next(shape for shape in check.slides[number - 1].shapes if shape.has_text_frame and shape.text == "Demo Index")
        assert link.click_action.target_slide == check.slides[24]
    assert sum(any(shape.has_text_frame and shape.text == "DEMO" for shape in slide.shapes)
               for slide in list(check.slides)[:18]) == 5
    for key, title, _ in destinations:
        number = int(key[-1])
        shape = next(shape for shape in check.slides[24].shapes if shape.has_text_frame and shape.text == f"{number:02d}  {title}")
        assert shape.click_action.target_slide == check.slides[MAIN_SEQUENCE.index(key)]
    print(f"PASS: 18 main slides, 6 hidden appendices, hidden Demo Index; architecture, demo breaks, notes and navigation verified: {OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build()